"""Exercise the frozen distribution using isolated data and preferences."""
from __future__ import annotations

import copy
import json
import sys
import time
import traceback
from pathlib import Path
from unittest.mock import patch


def run(destination: str) -> int:
    output = Path(destination).resolve()
    output.mkdir(parents=True, exist_ok=True)
    result = {"ok": False, "frozen": bool(getattr(sys, "frozen", False)), "checks": []}
    app = window = None

    def check(condition, label):
        if not condition:
            raise AssertionError(label)
        result["checks"].append(label)

    try:
        from PySide6.QtCore import QEvent, QPoint, QSettings, Qt, QUrl, qVersion
        from PySide6.QtGui import QColor, QImage, QImageReader, QKeyEvent, QTextDocument
        from PySide6.QtWidgets import QApplication, QDialog
        from openpyxl import load_workbook

        from detretlabel.annotation import AnnotationStore
        from detretlabel.app import MainWindow, STYLE
        from detretlabel.mapping_dialog import MappingDialog
        from detretlabel.review_report import write_excel
        from detretlabel.type_mapping import TypeMapping
        from detretlabel.app_info import AUTHOR, BUILD

        app = QApplication(["DetRetLabel"])
        app.setStyle("Fusion")
        app.setStyleSheet(STYLE)
        result["qt"] = qVersion()
        result["platform"] = app.platformName()
        bundle = Path(__file__).resolve().parent
        # In source runs this helper is in deployment/; in a bundle it is at its root.
        defaults = bundle / "config" / "reasons.json"
        if not getattr(sys, "frozen", False):
            defaults = bundle.parent / "config" / "reasons.json"
        check(defaults.is_file(), "Bundled default reason configuration")

        image = QImage(640, 400, QImage.Format.Format_RGB32)
        image.fill(QColor("#c7d9ed"))
        codecs = output / "图片格式"
        codecs.mkdir(exist_ok=True)
        for extension in ("jpg", "jpeg", "png", "bmp", "tif", "tiff", "webp"):
            path = codecs / f"sample.{extension}"
            check(image.save(str(path)) and not QImageReader(str(path)).read().isNull(), f"Image codec: {extension}")

        images = output / "测试 图片"
        images.mkdir(exist_ok=True)
        for filename in ("01.jpg", "02.png"):
            check(image.save(str(images / filename)), f"Create fixture {filename}")
        settings = QSettings(str(output / "settings.ini"), QSettings.Format.IniFormat)
        settings.clear()
        # Force use of the packaged default, independently of this machine's registry.
        settings.setValue("review/reasons_path", str(defaults))
        window = MainWindow(settings)
        window.show()
        app.processEvents()
        check(window.isVisible(), "Qt main window starts")
        check(bool(window.review.config.reasons["false_positive"]), "Default reasons load in the GUI")
        result["author"], result["build"] = AUTHOR, BUILD
        check(AUTHOR == "H9elix" and len(BUILD) == 13 and BUILD.startswith("Build"), "Author and build-time stamp")
        check(any(action.text() == f"打包：{BUILD}" for action in window.help_menu.actions()), "Build information in Help menu")
        window.readme_action.trigger()
        help_window = window._help_dialog
        app.processEvents()
        check(help_window.isVisible() and "快速开始" in help_window.tabs.widget(0).toPlainText(), "Bundled README renders in Help window")
        readme_browser = help_window.tabs.widget(0)
        help_window.open_link(readme_browser, QUrl("docs/使用指南.md"))
        check(help_window.tabs.currentIndex() == 1 and "审核算法结果" in help_window.tabs.currentWidget().toPlainText(), "Bundled user guide and relative document links")
        illustration = help_window.tabs.currentWidget().document().resource(
            QTextDocument.ResourceType.ImageResource, QUrl("images/05-gt-linkage.png"))
        check(illustration is not None and illustration.width() > 0, "Bundled illustrated guide loads offline images")
        help_window.search.setText("换机交接")
        help_window.search.setFocus()
        app.sendEvent(help_window.search, QKeyEvent(QEvent.Type.KeyPress, Qt.Key.Key_Return, Qt.KeyboardModifier.NoModifier))
        app.sendEvent(help_window.search, QKeyEvent(QEvent.Type.KeyRelease, Qt.Key.Key_Return, Qt.KeyboardModifier.NoModifier))
        check(help_window.isVisible() and help_window.tabs.currentWidget().textCursor().selectedText() == "换机交接", "Help document Enter searches without closing the window")
        check(help_window.grab().save(str(output / "help.png")), "Capture help preview")
        help_window.close()
        app.processEvents()

        def wait_idle():
            deadline = time.monotonic() + 20
            while window.busy and time.monotonic() < deadline:
                app.processEvents()
                time.sleep(0.01)
            app.processEvents()
            check(not window.busy, "Background operation completed")

        window.scan_folder(images)
        wait_idle()
        check(len(window.images) == 2, "Unicode and space-containing image folder")
        data = {"results": [
            {"path": "/tmp/01.jpg", "detections": [
                {"bbox": [220, 50, 300, 190], "type": "car", "conf": 0.8},
                {"bbox": [10, 20, 120, 180], "type": "person", "conf": 0.9}]},
            {"path": "C:\\other\\02.png", "detections": []}]}
        gt = copy.deepcopy(data)
        for record in gt["results"]:
            for target in record["detections"]:
                target.pop("conf", None)
                x1, y1, x2, y2 = target.pop("bbox")
                target.update(xmin=x1 / 640, ymin=y1 / 400, xmax=x2 / 640, ymax=y2 / 400)
        gt["results"][0]["detections"][0]["conf"] = .8
        for role, document in (("algorithm", data), ("ground_truth", gt)):
            path = output / f"{role}.json"
            path.write_text(json.dumps(document, ensure_ascii=False), encoding="utf-8")
            # Build the real mapping dialog, accepting its detected fields without user input.
            def accept_mapping(dialog):
                if role == "ground_truth":
                    dialog.normalized_coordinates.setChecked(True)
                    dialog.show()
                    app.processEvents()
                    check(all(combo.mapTo(dialog, QPoint(0, combo.height())).y() <=
                              dialog.category_combo.mapTo(dialog, QPoint()).y() for combo in dialog.coordinate_combos),
                          "Four-coordinate mapping fields do not overlap category or score")
                    dialog.close()
                return QDialog.DialogCode.Accepted
            with patch.object(MappingDialog, "exec", accept_mapping):
                window._map_document(document, path, role=role)
            wait_idle()
            check(window.sources[role].result.matched_target_count == 2, f"Import and map {role} JSON")

        gt_source = window.sources["ground_truth"]
        check(gt_source.mapping.normalized and gt_source.result.detections[images / "01.jpg"][0].box == (220, 50, 300, 190),
              "Normalized GT coordinates use matched image dimensions")
        check([item.full_text for item in window.image_view.label_items["ground_truth"]] == ["car", "person"],
              "GT labels hide confidence even when JSON contains scores")
        check(window.auto_gt_overlay.isChecked(), "Matched GT overlay enabled by default")

        mapping_path = output / "type-mapping.txt"
        mapping_path.write_text("person:行人\ncar:汽车\n", encoding="utf-8")
        window.set_type_mapping("algorithm", TypeMapping.load(mapping_path))
        check(window.type_mappings["algorithm"].display("person") == "行人", "Category display mapping")
        check(len(window.image_view.overlay_items["algorithm"]) == 2, "Detection box overlays")

        review = window.review
        check(review.current.category == "person" and review.current.target_index == 2,
              "Confidence ordering preserves original JSON identity")
        first_id = review.current.id
        check(window.annotation_panel.keep_reviewed_targets.isChecked(), "Retain reviewed targets enabled by default")
        review.choose_status("true_positive")
        view = window.image_view
        check(all(item.isVisible() for item in view.overlay_items["algorithm"]) and
              all(item.isVisible() for item in view.label_items["algorithm"]), "Reviewed box and label remain visible after advancing")
        window.annotation_panel.keep_reviewed_targets.setChecked(False)
        check(not view.overlay_items["algorithm"][0].isVisible() and not view.label_items["algorithm"][0].isVisible(),
              "Disabling retention hides noncurrent reviewed target")
        window.annotation_panel.keep_reviewed_targets.setChecked(True)
        check(all(item.isVisible() for item in view.label_items["algorithm"]), "Re-enabling retention restores reviewed labels")
        review.choose_status("false_positive")
        review.choose_reason("种类检错")
        check(review.phase == "manual" and review.algorithm_ready(), "03 decisions and automatic 04 transition")
        with patch("detretlabel.review_controller.QInputDialog.getText", return_value=("人工漏检", True)):
            review.add_manual((320, 100, 400, 250))
        review.choose_reason("完全未检出")
        check(len(review.store.manual_targets) == 1, "Manual missed detection in 04")
        check(review.save(explicit=True), "Save annotation progress")
        restored = AnnotationStore.load(images)
        check(restored.get(first_id)["status"] == "true_positive" and len(restored.manual_targets) == 1, "Reload saved progress")
        review.finish_review()
        check(review.review_completed(window.images[0]), "Complete 04 and retain image review status")
        report = review.report_snapshot()
        path = write_excel(report, output / "复核报告.xlsx")
        workbook = load_workbook(path)
        try:
            check(workbook.sheetnames == ["复核概览", "图片进度", "目标明细"], "Excel export and reopen: all three sheets")
            check(report.stats["true_positive"] == 1 and report.stats["false_positive"] == 1 and report.stats["false_negative"] == 1, "Exported review statistics")
        finally:
            workbook.close()
        window.image_table.setCurrentCell(0, 0)
        app.processEvents()
        check(window.grab().save(str(output / "window.png")), "Capture rendered GUI")
        result["ok"] = True
    except Exception:
        result["error"] = traceback.format_exc()
    finally:
        if window is not None:
            window.close()
        if app is not None:
            app.processEvents()
        (output / "result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if result["ok"] else 1
