"""Filterable Excel diagnostics for a target JSON export."""
from __future__ import annotations

import json
import os
from pathlib import Path
import re
import tempfile
from datetime import datetime, timedelta, timezone


ISSUE_NAMES = {
    "dimensions": "缺少可信图片尺寸", "size_conflict": "图片尺寸冲突", "size_mapping": "图片尺寸映射提示",
    "name": "图片路径无效", "type": "种类缺失或无效", "box": "目标框无效", "conf": "置信度无效",
    "precision": "小数精度问题", "type_name": "显示名称未匹配", "missing_conf": "置信度缺失",
    "outside_image": "目标框越界", "manual_type": "人工种类与算法类别不一致",
}


def _category_kind(value):
    if value is None:
        return "未指定"
    if isinstance(value, str):
        return "字符串"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "数值"
    return "无效类型"


def write_check_excel(export, path: Path, *, json_path: Path, allow_partial=False) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    if path.suffix.lower() != ".xlsx":
        raise ValueError("导出检查报告必须保存为 .xlsx。")
    blocked = bool(export.errors and not allow_partial)
    book = Workbook()
    temporary = None
    navy, blue = "24344B", "2563EB"

    def put(cell, value, header=False):
        if isinstance(value, str) and (len(value) > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value)):
            raise ValueError(f"导出检查报告 {cell.coordinate} 内容超过 Excel 限制或含无效控制字符。")
        cell.value = value
        if isinstance(value, str):
            cell.data_type = "s"  # Paths, categories and explanations are literal data.
        cell.font = Font(name="Microsoft YaHei", size=10, bold=header, color="FFFFFF" if header else navy)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if header:
            cell.fill = PatternFill("solid", fgColor=blue)

    try:
        details = book.active
        details.title = "检查明细"
        details.sheet_view.showGridLines = False
        headers = [("级别", 10), ("问题类型", 24), ("图片相对路径", 36), ("目标", 20),
                   ("目标唯一标识", 30), ("原始 type（JSON 值）", 26), ("type 值类型", 14),
                   ("type_name", 26), ("处理结果", 24), ("问题说明", 78)]
        for column, (label, width) in enumerate(headers, 1):
            put(details.cell(1, column), label, True)
            details.column_dimensions[get_column_letter(column)].width = width
        details.row_dimensions[1].height = 30
        issues = sorted(export.issues, key=lambda issue: issue.severity != "error")
        if len(issues) > 1048575:
            raise ValueError("检查明细超过 Excel 单表上限，未保存报告。")
        for row, issue in enumerate(issues, 2):
            state = "未生成 JSON" if blocked else "跳过该目标" if issue.severity == "error" else "保留并导出"
            category = "" if issue.category is None else json.dumps(issue.category, ensure_ascii=False)
            values = ("错误" if issue.severity == "error" else "提示", ISSUE_NAMES.get(issue.code, issue.code),
                issue.image, issue.target_label or "整图", issue.target_id, category,
                _category_kind(issue.category) if issue.target_id else "", issue.type_name, state, issue.message)
            for column, value in enumerate(values, 1):
                put(details.cell(row, column), value)
            lines = max(1 + sum(2 if ord(character) > 255 else 1 for character in str(value)) // int(width)
                        for value, (_, width) in zip(values, headers))
            details.row_dimensions[row].height = min(110, max(30, 16 * lines + 6))
        reference = f"A1:J{len(issues) + 1}"
        if issues:
            table = Table(displayName="ExportChecks", ref=reference)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            details.add_table(table)
        else:
            details.auto_filter.ref = reference
        details.freeze_panes = "A2"  # Only one header row; no large fixed summary region.

        summary = book.create_sheet("导出概览")
        summary.sheet_view.showGridLines = False
        summary.column_dimensions["A"].width = 24
        summary.column_dimensions["B"].width = 100
        for column, value in enumerate(("项目", "内容"), 1):
            put(summary.cell(1, column), value, True)
        summary.row_dimensions[1].height = 30
        status = "JSON 未生成，已有文件保留" if blocked else "仅导出有效部分" if export.errors else "导出完成"
        rows = [("导出结果", status), ("检查时间", datetime.now(timezone(timedelta(hours=8))).isoformat(timespec="seconds")),
            ("JSON 文件", str(json_path)), ("候选目标数", export.candidate_count),
            ("有效目标数", len(export.targets)), ("实际导出目标数", 0 if blocked else len(export.targets)),
            ("跳过目标数", 0 if blocked else export.candidate_count - len(export.targets)),
            ("错误项数", len(export.errors)), ("提示项数", len(export.warnings)),
            ("输出规范", "十字段平铺数组；归一化左上角 x / y / w / h；置信度和坐标八位普通十进制。"),
            ("筛选方法", "在“检查明细”首行使用筛选按钮，按级别、问题类型、图片、种类或处理结果筛选。"),
            ("种类说明", "原始 type 列按 JSON 值显示：数值 0 与字符串 \"0\" 不同；不自动转换旧人工类别。")]
        for row, (label, value) in enumerate(rows, 2):
            put(summary.cell(row, 1), label)
            put(summary.cell(row, 2), value)
            summary.row_dimensions[row].height = 32 if row < 11 else 46
        path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=path.parent, prefix="export-check-", suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        book.save(temporary)
        with temporary.open("r+b") as handle:
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        book.close()
        if temporary and temporary.exists():
            temporary.unlink()
    return path
