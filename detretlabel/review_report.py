"""Build a review snapshot and export a standalone, atomic Excel report."""
from __future__ import annotations

import math
import os
import re
import tempfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path

from .annotation import AnnotationStore, ReviewTarget
from .type_mapping import TypeMapping, category_text


@dataclass(frozen=True)
class ReviewReport:
    exported_at: str
    folder: str
    sources: dict[str, str]
    mappings: dict[str, str]
    stats: dict[str, int]
    images: list[dict]
    targets: list[dict]
    categories: list[dict]


def build_report(folder: Path, images: list[Path], by_image: dict[Path, list[ReviewTarget]],
                 store: AnnotationStore, signature: str, type_mappings: dict[str, TypeMapping],
                 sources: dict[str, str], input_states: dict[Path, str],
                 pending_reason: tuple[str, str] | None = None) -> ReviewReport:
    image_rows, target_rows = [], []
    stats = Counter(images=len(images), completed_images=0, algorithm=0, true_positive=0,
                    false_positive=0, unreviewed=0, false_negative=0, pending_manual=0,
                    pending_reasons=0, algorithm_completed_images=0, miss_completed_images=0)
    categories = defaultdict(Counter)
    pending_id, draft = pending_reason or (None, "")
    for path in images:
        image = path.relative_to(folder).as_posix()
        targets = by_image.get(path, [])
        counts = Counter(algorithm=0, true_positive=0, false_positive=0, unreviewed=0,
                         false_negative=0, pending_manual=0, pending_reasons=0)
        detail = []
        for role in ("algorithm", "manual"):
            for number, target in enumerate((t for t in targets if t.role == role), 1):
                record = store.get(target.id) or {}
                status = record.get("status")
                if role == "algorithm":
                    counts["algorithm"] += 1
                    bucket = status if status in ("true_positive", "false_positive") else "unreviewed"
                    status_text = {"true_positive": "正检", "false_positive": "误检", "false_negative": "旧漏检，待重审"}.get(status, "未审核")
                    if record.get("reason_pending"):
                        counts["pending_reasons"] += 1
                        status_text += "（原因待确认）"
                else:
                    bucket = "false_negative" if status == "false_negative" else "pending_manual"
                    status_text = "漏检" if bucket == "false_negative" else "待确认漏检"
                counts[bucket] += 1
                display = type_mappings.get("algorithm" if role == "manual" else role, TypeMapping()).display(target.category)
                raw = category_text(target.category)
                source = "算法结果" if role == "algorithm" else "人工补漏"
                categories[(source, raw, display if display is not None else "未指定种类")][bucket] += 1
                detail.append({"image": image, "source": source, "role": role, "status_code": status,
                    "reason_pending": bool(record.get("reason_pending")),
                    "category": target.category, "number": number, "raw_type": raw,
                    "display_type": display, "score": target.score, "xmin": target.box[0], "ymin": target.box[1],
                    "xmax": target.box[2], "ymax": target.box[3], "status": status_text,
                    "reason": record.get("reason", ""), "draft": draft if target.id == pending_id else record.get("reason_draft", ""),
                    "editing": "正在选择原因" if target.id == pending_id else "原因待确认" if record.get("reason_pending") else "",
                    "note": record.get("note", ""), "updated_at": record.get("updated_at", ""),
                    "target_id": target.id, "source_file": target.source_file})
        algorithm_completed = not counts["unreviewed"] and not counts["pending_reasons"]
        miss_completed = store.review_completed(image, signature) and not counts["pending_manual"]
        completed = algorithm_completed and miss_completed
        review_state = "复核完成" if completed else "03未完成 / 04已完成" if miss_completed else "待复核"
        image_review = store.image_reviews.get(image, {})
        image_rows.append({"image": image, **counts, "review_state": review_state,
            "completed_at": image_review.get("completed_at", "") if completed else "",
            "algorithm_state": "已完成" if algorithm_completed else "未完成",
            "miss_review_state": "已完成" if miss_completed else "待复核",
            "miss_completed_at": image_review.get("completed_at", "") if miss_completed else "",
            "input_state": input_states.get(path, ""),
            "gt_count": sum(t.role == "ground_truth" for t in targets)})
        for row in detail:
            row["review_state"] = review_state
            row["algorithm_state"] = image_rows[-1]["algorithm_state"]
            row["miss_review_state"] = image_rows[-1]["miss_review_state"]
        target_rows.extend(detail)
        stats.update(counts)
        stats["completed_images"] += int(completed)
        stats["algorithm_completed_images"] += int(algorithm_completed)
        stats["miss_completed_images"] += int(miss_completed)
    stats["pending_images"] = stats["images"] - stats["completed_images"]
    category_rows = [{"source": key[0], "raw_type": key[1], "display_type": key[2],
        **{name: value[name] for name in ("true_positive", "false_positive", "unreviewed", "false_negative", "pending_manual")}}
        for key, value in sorted(categories.items())]
    return ReviewReport(datetime.now(timezone(timedelta(hours=8))).isoformat(timespec="seconds"), str(folder),
        dict(sources), {role: f"{len(value.entries)} 条 · {value.path or '已保存配置'}" if value.entries else "未配置（显示原值）"
                       for role, value in type_mappings.items()}, dict(stats), image_rows, target_rows, category_rows)


def write_excel(report: ReviewReport, path: Path, *, targets_description: str = "") -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.hyperlink import Hyperlink
    except ImportError as exc:
        raise ValueError("缺少 Excel 导出依赖，请使用项目 Python 安装 requirements.txt 后重试。") from exc

    if path.suffix.lower() != ".xlsx":
        raise ValueError("复核报告必须保存为 .xlsx 文件。")
    book = Workbook()
    navy, blue = "24344B", "2563EB"
    plain_font = Font(name="Microsoft YaHei", size=10, color=navy)

    def put(cell, value):
        if isinstance(value, str):
            if len(value) > 32767 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
                raise ValueError(f"{cell.parent.title} {cell.coordinate} 内容过长或含 Excel 不支持的控制字符，无法完整导出。")
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError("报告含非有限数值，无法导出。")
        cell.value = value
        if isinstance(value, str):
            # Category codes, reasons and notes are literal data, including leading '='.
            cell.data_type = "s"
        cell.font = plain_font
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    def title(sheet, text, columns):
        sheet.sheet_view.showGridLines = False
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
        put(sheet.cell(1, 1), text)
        sheet.cell(1, 1).font = Font(name="Microsoft YaHei", size=18, bold=True, color="FFFFFF")
        sheet.cell(1, 1).fill = PatternFill("solid", fgColor=navy)
        sheet.row_dimensions[1].height = 36
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
        put(sheet.cell(2, 1), f"导出时间：{report.exported_at}  |  当前进度快照（包含未审核 / 待复核记录）")
        sheet.row_dimensions[2].height = 27
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0

    def table(sheet, row, columns, rows, name, *, freeze=True, grow_widths=False):
        if len(rows) + row > 1048576:
            raise ValueError(f"{sheet.title} 的记录数超过 Excel 单表上限。")
        for col, (key, label, width) in enumerate(columns, 1):
            letter = get_column_letter(col)
            previous = sheet.column_dimensions[letter].width if letter in sheet.column_dimensions else 0
            sheet.column_dimensions[letter].width = max(previous, width) if grow_widths else width
            cell = sheet.cell(row, col)
            put(cell, label)
            cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=blue)
        sheet.row_dimensions[row].height = 28
        for line, record in enumerate(rows, row + 1):
            height = 22
            for col, (key, _, width) in enumerate(columns, 1):
                value = record.get(key)
                cell = sheet.cell(line, col)
                put(cell, value)
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    cell.number_format = "0.####" if key in ("score", "xmin", "ymin", "xmax", "ymax") else "#,##0"
                elif key in ("raw_type", "display_type", "target_id"):
                    cell.number_format = "@"
                if isinstance(value, str):
                    height = max(height, min(80, 15 * (1 + len(value) // max(1, int(width))) + 5))
            sheet.row_dimensions[line].height = height
        end = f"{get_column_letter(len(columns))}{row + len(rows)}"
        if rows:
            data_table = Table(displayName=name, ref=f"A{row}:{end}")
            data_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(data_table)
        else:
            sheet.auto_filter.ref = f"A{row}:{end}"
        if freeze:
            sheet.freeze_panes = f"D{row + 1}" if len(columns) > 8 else f"A{row + 1}"
        sheet.print_title_rows = f"1:{row}"

    summary = book.active
    summary.title = "复核概览"
    title(summary, "检测结果复核报告", 8)
    metadata = [("报告状态", "全部复核完成" if report.images and not report.stats["pending_images"] else "复核进行中"),
        ("图片目录", report.folder), ("算法 JSON", report.sources.get("algorithm", "未导入")),
        ("GT JSON（仅参考）", report.sources.get("ground_truth", "未导入")),
        ("算法种类映射", report.mappings.get("algorithm", "未配置")), ("GT 种类映射", report.mappings.get("ground_truth", "未配置")),
        ("导出范围", "全部已加载图片；不受图片搜索筛选影响；不包含旧版本算法目标和 GT 人工历史标记。"),
        ("统计说明", "正检 / 误检来自算法目标人工判定，漏检来自人工补框。03 与 04 分别记录，仅两步都完成才算图片完成；原因待确认仍算 03 未完成，草稿不视为已确认原因。")]
    if targets_description:
        metadata.append(("有效目标 JSON", targets_description))
    for row, (label, value) in enumerate(metadata, 4):
        put(summary.cell(row, 1), label)
        summary.cell(row, 1).font = Font(name="Microsoft YaHei", size=10, bold=True, color=navy)
        summary.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        put(summary.cell(row, 2), value)
        summary.row_dimensions[row].height = 30 if row < 10 else 43
    metrics = [("图片总数", "images"), ("复核完成", "completed_images"), ("待复核图片", "pending_images"),
        ("算法目标总数", "algorithm"), ("正检", "true_positive"), ("误检", "false_positive"),
        ("算法未审核", "unreviewed"), ("人工漏检", "false_negative"), ("补框待确认", "pending_manual"),
        ("03 已完成图片", "algorithm_completed_images"), ("04 已完成图片", "miss_completed_images"), ("算法原因待确认", "pending_reasons")]
    for index, (label, key) in enumerate(metrics):
        row, col = 13 + index // 3, 1 + 3 * (index % 3)
        put(summary.cell(row, col), label)
        put(summary.cell(row, col + 1), report.stats[key])
        summary.cell(row, col + 1).font = Font(name="Microsoft YaHei", size=14, bold=True, color=blue)
        summary.row_dimensions[row].height = 28
    put(summary.cell(18, 1), "按来源及原始种类统计")
    category_columns = [("source", "来源", 18), ("raw_type", "原始种类", 20), ("display_type", "显示种类", 24),
        ("true_positive", "正检", 12), ("false_positive", "误检", 12), ("unreviewed", "未审核", 14),
        ("false_negative", "人工漏检", 16), ("pending_manual", "补框待确认", 16)]
    table(summary, 19, category_columns, report.categories, "ReviewCategories", freeze=False)
    progress = book.create_sheet("图片进度")
    title(progress, "图片复核进度", 15)
    table(progress, 4, [("image", "图片相对路径", 36), ("algorithm", "算法目标", 12), ("true_positive", "正检", 10),
        ("false_positive", "误检", 10), ("unreviewed", "算法未审核", 14), ("false_negative", "人工漏检", 12),
        ("pending_manual", "补框待确认", 14), ("review_state", "图片整体状态", 26),
        ("completed_at", "图片全部完成时间", 30), ("input_state", "算法输入状态", 24), ("gt_count", "GT参考框数", 14),
        ("algorithm_state", "03 算法审核", 18), ("pending_reasons", "算法原因待确认", 20),
        ("miss_review_state", "04 漏检复核", 18), ("miss_completed_at", "04 完成时间", 30)],
        report.images, "ImageReviewProgress")
    details = book.create_sheet("目标明细")
    detail_columns = [("image", "图片相对路径", 32), ("source", "目标来源", 14), ("number", "目标序号", 11),
        ("raw_type", "原始种类", 18), ("display_type", "显示种类", 22), ("score", "置信度（原始尺度）", 21),
        ("xmin", "xmin", 12), ("ymin", "ymin", 12), ("xmax", "xmax", 12), ("ymax", "ymax", 12),
        ("status", "人工结论", 20), ("reason", "原因（已保存）", 28), ("draft", "原因草稿（未确认）", 28),
        ("editing", "当前编辑状态", 20), ("note", "备注", 42), ("updated_at", "标注更新时间", 30),
        ("review_state", "图片整体状态", 26), ("target_id", "目标唯一标识", 24), ("source_file", "来源JSON", 40),
        ("algorithm_state", "03 算法审核", 18), ("miss_review_state", "04 漏检复核", 18)]
    title(details, "目标审核明细", len(detail_columns))
    table(details, 4, detail_columns, report.targets, "TargetReviewDetails")

    for label, status, role, prefix in (("正检", "true_positive", "algorithm", "TruePositive"),
                                       ("误检", "false_positive", "algorithm", "FalsePositive"),
                                       ("漏检", "false_negative", "manual", "FalseNegative")):
        selected = []
        grouped = {}
        for record in report.targets:
            if record["role"] != role or record["status_code"] != status:
                continue
            value = record["category"]
            kind = "未指定" if value is None else "字符串" if isinstance(value, str) else "数值"
            row = {**record, "type_value_kind": kind}
            selected.append(row)
            key = (kind, record["raw_type"], record["display_type"])
            category = grouped.setdefault(key, {"raw_type": record["raw_type"] or "未指定种类",
                "type_value_kind": kind, "display_type": record["display_type"] or "未指定种类",
                "count": 0, "images": set(), "pending_reasons": 0})
            category["count"] += 1
            category["images"].add(record["image"])
            category["pending_reasons"] += int(record.get("reason_pending", False))
        category_rows = [{**category, "image_count": len(category["images"])} for category in grouped.values()]
        category_rows.sort(key=lambda row: (-row["count"], row["type_value_kind"], row["raw_type"], row["display_type"]))
        sheet = book.create_sheet(label)
        page_columns = detail_columns[:4] + [("type_value_kind", "type 值类型", 14)] + detail_columns[4:]
        title(sheet, f"{label}结果概览与目标列表", len(page_columns))
        scope = "算法目标已判正检" if status == "true_positive" else "算法目标已判误检，原因待确认也列入并标明" if status == "false_positive" else "人工补框已标为漏检"
        put(sheet.cell(2, 1), f"导出时间：{report.exported_at}  |  {scope}；包含未完成图片中的已有结论。")
        for column, text, value in ((1, f"{label}目标数", len(selected)),
                                    (4, "涉及图片数", len({row['image'] for row in selected})),
                                    (7, "种类数", len(category_rows))):
            put(sheet.cell(4, column), text)
            put(sheet.cell(4, column + 1), value)
            sheet.cell(4, column + 1).font = Font(name="Microsoft YaHei", size=14, bold=True, color=blue)
        sheet.row_dimensions[4].height = 28
        put(sheet.cell(5, 1), "快速查看")
        sheet.merge_cells("B5:C5")
        put(sheet.cell(5, 2), "跳转到目标列表 ↓")
        if status == "false_positive":
            put(sheet.cell(5, 4), "原因待确认")
            put(sheet.cell(5, 5), sum(row.get("reason_pending", False) for row in selected))
        put(sheet.cell(5, 7), "未完成图片内目标")
        put(sheet.cell(5, 8), sum(row["review_state"] != "复核完成" for row in selected))
        sheet.row_dimensions[5].height = 28
        sheet.merge_cells("A6:H6")
        put(sheet.cell(6, 1), "按原始种类及数值 / 字符串类型分别统计；同名显示种类不会合并不同 type。未审核、待确认补框、GT 和旧算法漏检不计入本页。")
        sheet.row_dimensions[6].height = 34
        put(sheet.cell(7, 1), f"各个种类的{label}统计")
        type_columns = [("raw_type", "原始种类", 24), ("type_value_kind", "type 值类型", 14),
                        ("display_type", "显示种类", 26), ("count", f"{label}数量", 14),
                        ("image_count", "涉及图片数", 16)]
        if status == "false_positive":
            type_columns.append(("pending_reasons", "其中原因待确认", 20))
        table(sheet, 8, type_columns, category_rows, prefix + "Categories", freeze=False, grow_widths=True)
        if not selected:
            put(sheet.cell(9, 1), f"暂无{label}目标")
        list_title_row = 10 + max(1, len(category_rows))
        header_row = list_title_row + 1
        put(sheet.cell(list_title_row, 1), f"{label}目标列表（{len(selected)} 个）")
        table(sheet, header_row, page_columns, selected, prefix + "Targets", freeze=False, grow_widths=True)
        sheet.cell(5, 2).hyperlink = Hyperlink(ref="B5", location=f"'{label}'!A{header_row}", display="跳转到目标列表 ↓")
        sheet.cell(5, 2).font = Font(name="Microsoft YaHei", size=10, color=blue, underline="single")
        sheet.print_title_rows = f"{header_row}:{header_row}"
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=path.parent, prefix="review-report-", suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        book.save(temporary)
        with temporary.open("r+b") as handle:
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        book.close()
        if temporary and temporary.exists():
            temporary.unlink()
    return path
