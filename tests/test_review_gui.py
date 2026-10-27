import os
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import copy
import json
import shutil
import tempfile
import time
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from PySide6.QtCore import QEvent, QPoint, QPointF, QSettings, Qt
from PySide6.QtGui import QColor, QFontDatabase, QImage, QKeyEvent
from PySide6.QtTest import QSignalSpy, QTest
from PySide6.QtWidgets import QApplication, QDialog, QMessageBox

from detretlabel.app import MainWindow, STYLE
from detretlabel.annotation import AnnotationStore, ReviewConfig, ReviewTarget, source_signature
from detretlabel.core import Collection, guess_mapping, parse_detection
from detretlabel.mapping_dialog import MappingDialog
from detretlabel.type_mapping import TypeMapping
from detretlabel.box_style import BoxStyle, MANUAL_BOX_COLOR, mismatch_color


class ReviewGuiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])
        if not QFontDatabase.families() and os.name == "nt":
            QFontDatabase.addApplicationFont("C:/Windows/Fonts/msyh.ttc")
            QFontDatabase.addApplicationFont("C:/Windows/Fonts/segoeui.ttf")
        cls.app.setStyle("Fusion")
        cls.app.setStyleSheet(STYLE)

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name).resolve()
        image = QImage(640, 400, QImage.Format.Format_RGB32)
        image.fill(QColor("#c7d9ed"))
        for filename in ("01.jpg", "02.png", "03.jpg"):
            self.assertTrue(image.save(str(self.folder / filename)))
        self.settings = QSettings(str(self.folder / "settings.ini"), QSettings.Format.IniFormat)
        self.w = MainWindow(self.settings)
        self.w.resize(1440, 960)
        self.w.show()
        self.data = {"results": [
            {"path": "/tmp/01.jpg", "detections": [
                {"bbox": [10, 20, 120, 180], "type": "person", "conf": .9},
                {"bbox": [220, 50, 300, 190], "type": "car", "conf": .8}]},
            {"path": "/tmp/02.png", "detections": [{"bbox": [30, 40, 100, 200], "type": "cat", "conf": .7}]},
            {"path": "/tmp/03.jpg", "detections": []}]}
        self.w.scan_folder(self.folder)
        self.wait()
        self.import_source(self.data)

    def tearDown(self):
        self.wait()
        self.w.close()
        self.app.processEvents()
        self.temp.cleanup()

    def wait(self):
        deadline = time.monotonic() + 10
        while self.w.busy and time.monotonic() < deadline:
            self.app.processEvents()
            QTest.qWait(5)
        self.assertFalse(self.w.busy)
        self.app.processEvents()

    def import_source(self, data, role="algorithm"):
        path = self.folder / f"{role}.json"
        path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        with patch.object(MappingDialog, "exec", return_value=QDialog.DialogCode.Accepted):
            self.w._map_document(data, path, role=role)
        self.wait()

    def key(self, key, modifiers=Qt.KeyboardModifier.NoModifier):
        self.w.image_view.setFocus()
        QTest.keyClick(self.w.image_view, key, modifiers)
        self.app.processEvents()

    def matching_gt(self):
        self.import_source([{"name": "01.jpg", "objects": [
            {"bbox": [10, 20, 120, 180], "type": "car"},
            {"bbox": [12, 22, 118, 178], "type": "person"},
            {"bbox": [50, 50, 150, 200], "type": "person"},
            {"bbox": [400, 300, 500, 390], "type": "dog"}]}], "ground_truth")

    def wait_matches(self):
        reference = self.w.review.gt_reference
        deadline = time.monotonic() + 5
        while (reference.timer.isActive() or reference._job is not None or reference._pending is not None) and time.monotonic() < deadline:
            self.app.processEvents()
            QTest.qWait(5)
        self.app.processEvents()
        self.assertIsNone(reference._job)
        self.assertIsNone(reference._pending)

    def test_confidence_order_keeps_review_tables_overlays_and_report_synchronized(self):
        data = copy.deepcopy(self.data)
        data["results"][0]["detections"] = [
            {"bbox": [10, 20, 120, 180], "type": "low"},
            {"bbox": [220, 50, 300, 190], "type": "high-first", "conf": .9},
            {"bbox": [330, 50, 400, 190], "type": "high-second", "conf": .9}]
        original = copy.deepcopy(data)
        self.import_source(data)
        review, view = self.w.review, self.w.image_view
        table = self.w.target_tables["algorithm"]
        expected = ["high-first", "high-second", "low"]
        self.assertEqual([table.item(row, 1).text() for row in range(3)], expected)
        self.assertEqual(table.item(2, 2).text(), "—")
        self.assertEqual([target.category for target in review.queue()], expected)
        self.assertEqual([target.target_index for target in review.queue()], [2, 3, 1])
        self.assertEqual(review.current.category, "high-first")
        for index, target in enumerate(review.queue()):
            item = view.overlay_items["algorithm"][index]
            self.assertEqual(item.rect().left(), target.box[0])
            self.assertEqual(item.data(1), index + 1)
        review.clicked_target("algorithm", 2)
        self.assertEqual((review.current.category, table.currentRow(), view.active_index), ("high-second", 1, 2))
        table.setCurrentCell(2, 0)
        self.assertEqual((review.current.category, review.current.target_index, view.active_index), ("low", 1, 3))
        self.assertIn("第 1 个目标", table.item(2, 0).toolTip())
        low_id = review.current.id
        review.choose_status("true_positive")
        self.assertEqual((review.current.category, table.currentRow(), view.active_index), ("high-first", 0, 1))
        self.assertEqual(review.store.get(low_id)["status"], "true_positive")
        details = [row for row in review.report_snapshot().targets if row["image"] == "01.jpg"]
        self.assertEqual([row["raw_type"] for row in details], expected)
        self.assertEqual([row["number"] for row in details], [1, 2, 3])
        cached = self.w.sources["algorithm"].result.detections[self.folder / "01.jpg"]
        with patch("detretlabel.core.sort_detections", side_effect=AssertionError("Must not sort while switching images")):
            self.w.move_image(1)
            self.w.move_image(-1)
        self.assertIs(self.w.sources["algorithm"].result.detections[self.folder / "01.jpg"], cached)
        self.assertEqual(review.current.category, "high-first")
        self.assertEqual(data, original)

    def test_sorted_gt_rows_clicks_and_iou_hints_use_matching_overlay_indices(self):
        gt = [{"name": "01.jpg", "objects": [
            {"bbox": [10, 20, 120, 180], "type": "car", "conf": .1},
            {"bbox": [12, 22, 118, 178], "type": "person", "conf": .9},
            {"bbox": [10, 20, 120, 180], "type": "person", "conf": .8},
            {"bbox": [450, 300, 500, 350], "type": "unscored"}]}]
        self.import_source(gt, "ground_truth")
        self.wait_matches()
        review, view = self.w.review, self.w.image_view
        table = self.w.target_tables["ground_truth"]
        self.assertEqual([float(table.item(row, 2).text()) for row in range(3)], [.9, .8, .1])
        self.assertEqual(table.item(3, 2).text(), "—")
        self.assertEqual([target.target_index for target in review.role_targets("ground_truth")], [2, 3, 1, 4])
        result = review.gt_reference.result
        self.assertEqual((result.same_index, result.best_index, result.mismatch_index), (2, 3, 3))
        self.assertEqual(table.item(1, 0).background().color().name(), "#e7f7ec")
        self.assertEqual(table.item(2, 0).background().color().name(), "#fff5d6")
        self.w.auto_gt_overlay.setChecked(True)
        self.assertEqual((view.same_gt_index, view.mismatch_gt_index), (2, 3))
        target_id = review.current.id
        table.setCurrentCell(0, 0)
        self.assertEqual(view.selected_gt_index, 1)
        self.assertEqual(view.overlay_items["ground_truth"][0].rect().left(), 12)
        review.clicked_target("ground_truth", 3)
        self.assertEqual((table.currentRow(), view.selected_gt_index), (2, 3))
        self.assertEqual(view.overlay_items["ground_truth"][2].rect().left(), 10)
        self.assertEqual(review.current.id, target_id)
        self.assertEqual(review.phase, "algorithm")

    def test_annotations_created_before_sorting_keep_their_target_after_reopen(self):
        data = copy.deepcopy(self.data)
        data["results"][0]["detections"][0]["conf"] = .2
        mapping = guess_mapping(data, Collection("/results"))
        old_detection = parse_detection(data["results"][0]["detections"][0], mapping, 1, 1)
        legacy = ReviewTarget.from_detection("01.jpg", "algorithm", self.folder / "algorithm.json",
                                             source_signature(data, mapping), old_detection)
        self.w.review.store.set_annotation(legacy, "false_positive", "原有原因", "原有备注")
        self.assertTrue(self.w.review.save(explicit=True))
        self.import_source(data)
        queue = self.w.review.queue()
        self.assertEqual([target.category for target in queue], ["car", "person"])
        self.assertEqual(queue[1].id, legacy.id)
        self.assertEqual(self.w.target_tables["algorithm"].item(1, 7).text(), "误检")
        self.w.close()
        self.w = MainWindow(self.settings)
        self.w.show()
        self.w.scan_folder(self.folder)
        self.wait()
        self.import_source(data)
        self.assertEqual(self.w.review.queue()[1].id, legacy.id)
        self.assertEqual(self.w.review.current.category, "car")
        record = self.w.review.store.get(legacy.id)
        self.assertEqual((record["status"], record["reason"], record["note"]), ("false_positive", "原有原因", "原有备注"))

    def test_tables_are_side_by_side_and_algorithm_click_returns_to_review(self):
        self.matching_gt()
        left = self.w.target_tables["algorithm"]
        right = self.w.target_tables["ground_truth"]
        self.assertTrue(left.isVisible())
        self.assertTrue(right.isVisible())
        self.assertGreater(right.mapToGlobal(QPoint()).x(), left.mapToGlobal(QPoint()).x())
        self.w.review.all_positive()
        self.assertEqual(self.w.review.phase, "manual")
        self.w.target_tabs.setCurrentWidget(left)
        QTest.mouseClick(left.viewport(), Qt.MouseButton.LeftButton, pos=left.visualItemRect(left.item(0, 0)).center())
        self.assertEqual(self.w.review.phase, "algorithm")
        self.assertEqual(self.w.review.current.category, "person")
        self.assertEqual(self.w.target_tabs.currentWidget(), left)

    def test_result_tables_get_vertical_growth_and_match_hint_stays_one_line(self):
        self.w.resize(1900, 1100)
        parent = self.w.target_splitter.parentWidget()
        vertical = parent.parentWidget()
        vertical.setSizes([650, 200])
        self.app.processEvents()
        initial = (parent.height(), self.w.target_splitter.height())
        vertical.setSizes([300, 550])
        self.app.processEvents()
        hint = self.w.gt_match_status
        self.assertLessEqual(hint.height(), hint.fontMetrics().height() + 4)
        self.assertGreater(self.w.target_splitter.height(), initial[1] + 100)
        self.assertAlmostEqual(parent.height() - initial[0], self.w.target_splitter.height() - initial[1], delta=3)
        self.assertLessEqual(hint.y() - (self.w.target_splitter.y() + self.w.target_splitter.height()), 4)
        self.assertLessEqual(parent.height() - (hint.y() + hint.height()), 2)

    def test_manual_reference_tab_and_rows_preserve_algorithm_review_and_reason_draft(self):
        self.matching_gt()
        review, panel, view = self.w.review, self.w.annotation_panel, self.w.image_view
        targets = [ReviewTarget.manual("01.jpg", (300 + i * 80, 10, 350 + i * 80, 80), category)
                   for i, category in enumerate(("dog", "cat"))]
        for target in targets:
            review.store.set_annotation(target, "false_negative", "人工确认", "补框备注")
        review.rebuild()
        review.fill_manual_table()
        self.w._update_overlays()
        panel.auto_next.setChecked(False)
        review.choose_status("false_positive")
        panel.reason_combo.setCurrentText("尚未确认的原因")
        current = review.current
        before = copy.deepcopy(review.store.records)
        table = self.w.target_tables["manual"]
        self.assertEqual(self.w.target_tabs.indexOf(table), -1)
        self.assertEqual(self.w.gt_tabs.indexOf(table), 1)
        self.w.gt_tabs.setCurrentWidget(table)
        table.setCurrentCell(1, 0)
        self.assertEqual(view.selected_manual_index, 2)
        self.assertEqual(view.overlay_items["manual"][1].pen().widthF(), 3)
        self.assertEqual(review.current, current)
        self.assertEqual(review.phase, "algorithm")
        self.assertEqual(review.reason_status, "false_positive")
        self.assertEqual(panel.reason_combo.currentText(), "尚未确认的原因")
        self.assertEqual(review.store.records, before)
        review.clicked_target("manual", 1)
        self.assertEqual(view.selected_manual_index, 1)
        self.assertEqual(review.current, current)
        self.w.gt_tabs.setCurrentWidget(self.w.target_tables["ground_truth"])
        self.w.target_tables["ground_truth"].setCurrentCell(3, 0)
        self.assertEqual(view.selected_gt_index, 4)
        self.assertEqual(view.selected_manual_index, 0)
        self.w.gt_tabs.setCurrentWidget(table)
        self.assertEqual(view.selected_manual_index, 1)
        self.assertEqual(view.selected_gt_index, 0)
        self.assertEqual(review.current, current)
        self.assertEqual(review.store.records, before)

    def test_manual_style_is_purple_and_tracks_gt_style_without_resetting_zoom(self):
        review, view = self.w.review, self.w.image_view
        target = ReviewTarget.manual("01.jpg", (300, 10, 350, 80), "dog")
        review.store.set_annotation(target, "false_negative")
        review.rebuild()
        review.fill_manual_table()
        self.w._update_overlays()
        self.assertEqual(view.overlay_items["manual"][0].pen().style(), Qt.PenStyle.DashLine)
        self.assertEqual(view.overlay_items["manual"][0].pen().color().name(), "#b77aff")
        view.auto_fit = False
        view.scale(1.5, 1.5)
        zoom = view.transform()
        style = BoxStyle("#125678", "dash_dot_dot", show_category=False, show_source=True)
        with patch("detretlabel.app.BoxStyleDialog") as dialog:
            dialog.return_value.exec.return_value = QDialog.DialogCode.Accepted
            dialog.return_value.selected_style.return_value = style
            self.w.configure_box_style("ground_truth")
        self.assertEqual(view.overlay_items["manual"][0].pen().style(), Qt.PenStyle.DashDotDotLine)
        self.assertEqual(view.overlay_items["manual"][0].pen().color().name(), "#b77aff")
        self.assertEqual(view.label_items["manual"][0].full_text, "人工补框")
        self.assertEqual(view.transform(), zoom)

    def test_manual_single_click_inspects_and_double_click_edits_existing_miss(self):
        review = self.w.review
        review.all_positive()
        with patch("detretlabel.review_controller.QInputDialog.getText", return_value=("dog", True)):
            review.add_manual((300, 10, 350, 80))
        target_id = review.current.id
        review.choose_reason("原原因")
        review.leave_review()
        algorithm_id = review.current.id
        review.store.set_image_review("01.jpg", review.algorithm_signature(), True)
        review.changed()
        table = self.w.target_tables["manual"]
        self.w.gt_tabs.setCurrentWidget(table)
        self.app.processEvents()
        point = table.visualItemRect(table.item(0, 0)).center()
        QTest.mouseClick(table.viewport(), Qt.MouseButton.LeftButton, pos=point)
        self.assertEqual(review.current.id, algorithm_id)
        self.assertEqual(review.phase, "algorithm")
        self.assertTrue(review.review_completed())
        QTest.mouseDClick(table.viewport(), Qt.MouseButton.LeftButton, pos=point)
        self.app.processEvents()
        self.assertEqual(review.current.id, target_id)
        self.assertEqual(review.phase, "manual")
        self.assertEqual(review.reason_status, "false_negative")
        self.assertFalse(review.review_completed())
        review.choose_reason("修正原因")
        self.assertEqual(review.store.get(target_id)["reason"], "修正原因")

    def test_matching_highlights_both_candidates_and_overlay_is_optional(self):
        self.matching_gt()
        self.wait_matches()
        review, view = self.w.review, self.w.image_view
        reference = review.gt_reference
        self.assertEqual((reference.result.same_index, reference.result.mismatch_index), (2, 1))
        table = self.w.target_tables["ground_truth"]
        self.assertEqual(table.item(1, 1).background().color().name(), "#e7f7ec")
        self.assertEqual(table.item(0, 1).background().color().name(), "#fff5d6")
        self.assertFalse(table.selectedItems())
        self.assertIn("IoU", self.w.gt_match_status.full_text)
        self.assertEqual(view.overlay_items["algorithm"][0].pen().color().name(), MANUAL_BOX_COLOR)
        self.assertTrue(self.w.auto_gt_overlay.isChecked())
        self.assertEqual([item.isVisible() for item in view.overlay_items["ground_truth"]], [True, True, False, False])
        self.w.auto_gt_overlay.setChecked(False)
        self.assertFalse(any(item.isVisible() for item in view.overlay_items["ground_truth"]))
        target_id = review.current.id
        self.w.auto_gt_overlay.setChecked(True)
        gt = view.overlay_items["ground_truth"]
        self.assertEqual([item.isVisible() for item in gt], [True, True, False, False])
        self.assertEqual(gt[0].pen().color().name(), "#ffff00")
        self.assertEqual(gt[1].pen().color().name(), "#00ff00")
        self.assertEqual(gt[0].pen().style(), Qt.PenStyle.DashLine)
        self.assertEqual(view.label_items["ground_truth"][0].color.name(), "#ffff00")
        self.w.annotation_panel.gt_button.setChecked(True)
        self.assertTrue(all(item.isVisible() for item in gt))
        self.w.annotation_panel.gt_button.setChecked(False)
        self.assertEqual([item.isVisible() for item in gt], [True, True, False, False])
        self.w.auto_gt_overlay.setChecked(False)
        self.assertFalse(any(item.isVisible() for item in gt))
        self.assertEqual(gt[0].pen().color().name(), "#00ff00")
        self.assertEqual(review.current.id, target_id)
        self.assertFalse(review.store.records)
        self.assertFalse(review.store.image_reviews)

    def test_manual_gt_selection_keeps_algorithm_target_reason_and_annotation(self):
        self.w.auto_gt_overlay.setChecked(False)
        self.matching_gt()
        self.wait_matches()
        review, panel, view = self.w.review, self.w.annotation_panel, self.w.image_view
        panel.auto_next.setChecked(False)
        review.choose_status("false_positive")
        panel.reason_combo.setCurrentText("尚未确认的原因")
        target = review.current
        before = copy.deepcopy(review.store.records)
        self.w.target_tables["ground_truth"].setCurrentCell(3, 0)
        self.assertEqual(review.current, target)
        self.assertEqual(review.store.records, before)
        self.assertEqual(review.reason_status, "false_positive")
        self.assertEqual(panel.reason_combo.currentText(), "尚未确认的原因")
        self.assertEqual([item.isVisible() for item in view.overlay_items["ground_truth"]], [False, False, False, True])
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [True, False])
        self.assertFalse(panel.gt_button.isChecked())
        review.move_target(1)
        self.assertEqual(review.current.category, "car")
        self.assertEqual(review.gt_reference.selected_index, 0)
        self.assertFalse(any(item.isVisible() for item in view.overlay_items["ground_truth"]))

    def test_reference_colors_follow_styles_and_avoid_source_colors(self):
        self.matching_gt()
        self.wait_matches()
        view = self.w.image_view
        view.auto_fit = False
        view.scale(1.6, 1.6)
        zoom = view.transform()
        self.w.auto_gt_overlay.setChecked(True)
        for role, style in (("algorithm", BoxStyle("#ffff00", "solid")), ("ground_truth", BoxStyle("#ff00ff", "dot"))):
            with patch("detretlabel.app.BoxStyleDialog") as dialog:
                dialog.return_value.exec.return_value = QDialog.DialogCode.Accepted
                dialog.return_value.selected_style.return_value = style
                self.w.configure_box_style(role)
        gt = view.overlay_items["ground_truth"]
        self.assertEqual(gt[1].pen().color().name(), "#ff00ff")
        self.assertNotIn(gt[0].pen().color().name(), ("#ffff00", "#ff00ff"))
        self.assertEqual(gt[0].pen().color().name(), mismatch_color("#ffff00", "#ff00ff"))
        self.assertEqual(gt[0].pen().style(), Qt.PenStyle.DotLine)
        self.assertEqual(gt[1].pen().style(), Qt.PenStyle.DotLine)
        self.assertEqual(view.transform(), zoom)

    def test_gt_matching_compares_mapped_names_and_does_not_change_review_state(self):
        self.matching_gt()
        coded = copy.deepcopy(self.data)
        coded["results"][0]["detections"][0]["type"] = 0
        self.import_source(coded)
        target = self.w.review.current
        self.w.review.store.set_annotation(target, "true_positive", note="已审核")
        self.w.set_type_mapping("algorithm", TypeMapping({"0": "行人"}))
        self.w.set_type_mapping("ground_truth", TypeMapping({"person": "行人", "car": "汽车"}))
        self.wait_matches()
        self.assertEqual(self.w.review.gt_reference.result.same_index, 2)
        self.assertEqual(self.w.review.gt_reference.result.mismatch_index, 1)
        self.assertEqual(self.w.review.store.get(target.id)["note"], "已审核")
        self.assertEqual(self.w.review.current.category, 0)

    def test_slow_gt_matching_does_not_block_image_change_or_apply_stale_results(self):
        from threading import Event
        from detretlabel.gt_matching import match_gt
        started, release = Event(), Event()
        def delayed(box, category, candidates, cancel):
            started.set()
            release.wait(3)
            return match_gt(box, category, candidates)  # Deliberately ignore cancellation to exercise stale-result handling.
        with patch("detretlabel.gt_reference.match_gt", side_effect=delayed):
            try:
                self.matching_gt()
                deadline = time.monotonic() + 2
                while not started.is_set() and time.monotonic() < deadline:
                    self.app.processEvents()
                    QTest.qWait(5)
                self.assertTrue(started.is_set())
                start = time.monotonic()
                self.w.image_table.setCurrentCell(1, 0)
                self.app.processEvents()
                self.assertEqual(self.w.current_image().name, "02.png")
                self.assertLess(time.monotonic() - start, .75)
                self.assertFalse(self.w.busy)
                self.assertIsNone(self.w.review.gt_reference.result)
            finally:
                release.set()
                self.wait_matches()
        self.assertIsNone(self.w.review.gt_reference.result)
        self.assertFalse(self.w.review.gt_reference._painted)
        self.assertEqual(self.w.review.current.category, "cat")
        self.assertIn("没有可匹配", self.w.gt_match_status.full_text)

    def test_rapid_target_changes_coalesce_to_the_latest_reference_request(self):
        from threading import Event
        from detretlabel.gt_matching import match_gt
        started, release = Event(), Event()
        calls = []
        def delayed(box, category, candidates, cancel):
            calls.append(category)
            started.set()
            release.wait(3)
            return match_gt(box, category, candidates)
        with patch("detretlabel.gt_reference.match_gt", side_effect=delayed):
            try:
                self.matching_gt()
                deadline = time.monotonic() + 2
                while not started.is_set() and time.monotonic() < deadline:
                    self.app.processEvents()
                    QTest.qWait(5)
                self.assertTrue(started.is_set())
                review = self.w.review
                review.move_target(1)
                review.move_target(-1)
                review.move_target(1)
                self.assertFalse(self.w.busy)
            finally:
                release.set()
                self.wait_matches()
        self.assertEqual(calls, ["person", "car"])
        self.assertEqual(self.w.review.gt_reference.result.best_index, 0)
        self.assertFalse(self.w.review.gt_reference._painted)

    def test_close_removes_global_listener_reopen_restores_it_and_gt_option_persists(self):
        from detretlabel.review_controller import ReviewController
        calls = []
        class TrackingReviewController(ReviewController):
            def eventFilter(controller, watched, event):
                if event.type() == QEvent.Type.User:
                    calls.append(watched)
                return super().eventFilter(watched, event)
        self.w.close()
        with patch("detretlabel.app.ReviewController", TrackingReviewController):
            self.w = MainWindow(self.settings)
        self.w.show()
        self.w.scan_folder(self.folder)
        self.wait()
        self.import_source(self.data)
        review = self.w.review
        first_id = review.current.id
        self.app.sendEvent(self.w.image_view, QEvent(QEvent.Type.User))
        self.assertTrue(calls)
        self.w.auto_gt_overlay.setChecked(True)
        self.w.close()
        calls.clear()
        self.app.sendEvent(self.w.image_view, QEvent(QEvent.Type.User))
        self.assertFalse(calls)
        self.w.show()
        self.app.sendEvent(self.w.image_view, QEvent(QEvent.Type.User))
        self.assertTrue(calls)
        self.key(Qt.Key.Key_1)
        self.assertEqual(review.store.get(first_id)["status"], "true_positive")
        self.w.close()
        self.w = MainWindow(self.settings)
        self.w.show()
        self.assertTrue(self.w.auto_gt_overlay.isChecked())

    def test_two_stage_numeric_shortcuts_and_auto_advance(self):
        review = self.w.review
        first = review.current.id
        self.key(Qt.Key.Key_2)
        self.assertEqual(review.store.get(first)["status"], "false_positive")
        self.assertEqual(review.reason_status, "false_positive")
        self.assertEqual(review.current.id, first)
        self.key(Qt.Key.Key_1)
        self.assertEqual(review.store.get(first)["reason"], "种类检错")
        second = review.current.id
        self.assertNotEqual(first, second)
        self.key(Qt.Key.Key_3)
        self.assertFalse(review.store.completed(second))
        self.key(Qt.Key.Key_1)
        self.assertEqual(review.store.get(second)["status"], "true_positive")
        self.assertIsNone(review.reason_status)
        self.assertTrue(self.w.image_table.item(0, 3).text().startswith("2/2"))
        self.assertIn("待复核", self.w.image_table.item(0, 3).text())

    def test_keypad_digits_choose_decisions_all_nine_reasons_and_enter_confirms(self):
        review, panel = self.w.review, self.w.annotation_panel
        reasons = tuple(f"reason {index}" for index in range(1, 10))
        review.apply_config(ReviewConfig(reasons={"false_positive": reasons, "false_negative": reasons}, auto_advance=False))
        keypad = Qt.KeyboardModifier.KeypadModifier
        target_id = review.current.id
        self.key(Qt.Key.Key_1, keypad)
        self.assertEqual(review.store.get(target_id)["status"], "true_positive")
        for index, reason in enumerate(reasons, 1):
            self.key(Qt.Key.Key_2, keypad)
            self.assertEqual(review.reason_status, "false_positive")
            self.key(Qt.Key(Qt.Key.Key_0.value + index), keypad)
            self.assertEqual(review.store.get(target_id)["reason"], reason)
            self.assertFalse(review.store.get(target_id).get("reason_pending", False))
        self.key(Qt.Key.Key_2, keypad)
        panel.reason_combo.setCurrentText("")
        self.key(Qt.Key.Key_Enter, keypad)
        self.assertIsNone(review.reason_status)
        self.assertEqual(review.store.get(target_id)["reason"], "")

    def test_keypad_digits_respect_text_entry_and_ignore_autorepeat(self):
        review, panel, view = self.w.review, self.w.annotation_panel, self.w.image_view
        keypad = Qt.KeyboardModifier.KeypadModifier
        first = review.current.id
        panel.note.setFocus()
        QTest.keyClick(panel.note, Qt.Key.Key_1, keypad)
        QTest.keyClick(panel.note, Qt.Key.Key_2, keypad)
        self.assertEqual(panel.note.text(), "12")
        self.assertFalse(review.store.completed(first))
        panel.auto_next.setChecked(True)
        self.key(Qt.Key.Key_1, keypad)
        second = review.current.id
        self.assertNotEqual(first, second)
        event = QKeyEvent(QEvent.Type.KeyPress, Qt.Key.Key_1, keypad, "1", True, 1)
        QApplication.sendEvent(view, event)
        self.assertFalse(review.store.completed(second))

    def test_explicit_keypad_binding_takes_priority_and_control_modifier_is_preserved(self):
        review = self.w.review
        config = ReviewConfig(auto_advance=False)
        config.shortcuts["decision_1"] = "Num+1"
        review.apply_config(config)
        keypad = Qt.KeyboardModifier.KeypadModifier
        target_id = review.current.id
        self.key(Qt.Key.Key_1)
        self.assertFalse(review.store.completed(target_id))
        self.key(Qt.Key.Key_1, keypad | Qt.KeyboardModifier.ShiftModifier)
        self.assertFalse(review.store.completed(target_id))
        self.key(Qt.Key.Key_1, keypad)
        self.assertEqual(review.store.get(target_id)["status"], "true_positive")
        self.key(Qt.Key.Key_1, keypad | Qt.KeyboardModifier.ControlModifier)
        self.assertTrue(review.algorithm_ready())
        self.assertEqual(review.phase, "manual")

    def test_mouse_buttons_reason_and_status_change_clear_reason(self):
        self.w.annotation_panel.auto_next.setChecked(False)
        panel, review = self.w.annotation_panel, self.w.review
        panel.wheel.buttons["false_positive"].click()
        panel.reason_buttons[1].click()
        self.assertEqual(review.store.get(review.current.id)["reason"], "单纯误检")
        panel.wheel.buttons["true_positive"].click()
        self.assertEqual(review.store.get(review.current.id)["status"], "true_positive")
        self.assertEqual(review.store.get(review.current.id)["reason"], "")
        self.assertEqual(self.w.target_table.item(0, 7).text(), "正检")

    def test_gt_is_reference_only_in_both_phases(self):
        gt = [{"name": "01.jpg", "objects": [{"xywh": [12, 22, 100, 150], "type": "GT person"}]}]
        self.import_source(gt, "ground_truth")
        review = self.w.review
        self.assertEqual(self.w.annotation_panel.scope_combo.currentData(), "algorithm")
        self.assertEqual(len(review.queue()), 2)
        self.assertEqual(review.current.role, "algorithm")
        self.w.target_tables["ground_truth"].setCurrentCell(0, 0)
        self.assertEqual(review.current.role, "algorithm")
        self.w.annotation_panel.scope_combo.setCurrentIndex(1)
        self.assertEqual(review.phase, "manual")
        self.assertFalse(review.algorithm_ready())
        self.assertFalse(review.store.records)
        review.leave_review()
        review.all_positive()
        self.w.annotation_panel.scope_combo.setCurrentIndex(1)
        self.assertEqual(review.phase, "manual")
        self.assertIsNone(review.current)
        self.assertEqual(len(review.queue()), 0)
        self.assertFalse(self.w.annotation_panel.gt_button.isChecked())
        self.assertEqual(self.w.target_tables["ground_truth"].item(0, 7).text(), "仅参考")

    def test_custom_reason_notes_and_reopen_restore(self):
        self.w.annotation_panel.auto_next.setChecked(False)
        review = self.w.review
        target_id = review.current.id
        self.key(Qt.Key.Key_2)
        panel = self.w.annotation_panel
        panel.reason_combo.setCurrentText("配置之外的原因")
        panel.confirm_button.click()
        panel.note.setFocus()
        QTest.keyClicks(panel.note, "123 note")
        self.assertTrue(review.save(explicit=True))
        self.assertEqual(review.store.get(target_id)["note"], "123 note")
        self.assertEqual(review.store.get(target_id)["reason"], "配置之外的原因")
        self.w.close()
        self.w = MainWindow(self.settings)
        self.w.show()
        self.w.scan_folder(self.folder)
        self.wait()
        self.import_source(self.data)
        self.assertEqual(self.w.review.store.get(target_id)["reason"], "配置之外的原因")
        self.assertEqual(self.w.target_table.item(0, 7).text(), "误检")

    def test_digits_in_note_do_not_trigger_review_and_note_survives_navigation(self):
        panel, review = self.w.annotation_panel, self.w.review
        target_id = review.current.id
        panel.note.setFocus()
        QTest.keyClicks(panel.note, "123")
        self.assertFalse(review.store.completed(target_id))
        self.w.image_table.setCurrentCell(1, 0)
        self.assertEqual(review.store.get(target_id)["note"], "123")
        self.assertIsNone(review.store.get(target_id)["status"])

    def test_navigation_undo_clear_and_next_unreviewed(self):
        review = self.w.review
        self.w.annotation_panel.auto_next.setChecked(False)
        first = review.current.id
        self.key(Qt.Key.Key_1)
        self.key(Qt.Key.Key_Down)
        second = review.current.id
        self.assertNotEqual(first, second)
        self.key(Qt.Key.Key_Up)
        self.assertEqual(review.current.id, first)
        self.key(Qt.Key.Key_Backspace)
        self.assertFalse(review.store.completed(first))
        self.key(Qt.Key.Key_Z, Qt.KeyboardModifier.ControlModifier)
        self.assertTrue(review.store.completed(first))
        self.key(Qt.Key.Key_N)
        self.assertEqual(review.current.id, second)
        self.key(Qt.Key.Key_1)
        self.key(Qt.Key.Key_N)
        self.assertEqual(self.w.current_image().name, "02.png")
        self.key(Qt.Key.Key_Left)
        self.assertEqual(self.w.current_image().name, "01.jpg")

    def test_source_changes_do_not_attach_old_annotations_to_different_data(self):
        self.w.annotation_panel.auto_next.setChecked(False)
        old_id = self.w.review.current.id
        self.key(Qt.Key.Key_1)
        changed = copy.deepcopy(self.data)
        changed["results"][0]["detections"][0]["bbox"][0] = 11
        self.import_source(changed)
        self.assertNotEqual(self.w.review.current.id, old_id)
        self.assertFalse(self.w.review.store.completed(self.w.review.current.id))
        self.import_source(self.data)
        self.assertTrue(self.w.review.store.completed(old_id))

    def test_manual_box_drag_creates_missed_target_and_delete_is_undoable(self):
        self.w.image_table.setCurrentCell(2, 0)
        self.key(Qt.Key.Key_R)
        self.key(Qt.Key.Key_B)
        self.assertTrue(self.w.image_view.draw_mode)
        view = self.w.image_view
        start, end = view.mapFromScene(QPointF(30, 40)), view.mapFromScene(QPointF(160, 180))
        with patch("detretlabel.review_controller.QInputDialog.getText", return_value=("dog", True)):
            QTest.mousePress(view.viewport(), Qt.MouseButton.LeftButton, pos=start)
            QTest.mouseMove(view.viewport(), end)
            QTest.mouseRelease(view.viewport(), Qt.MouseButton.LeftButton, pos=end)
        target = self.w.review.current
        self.assertEqual(target.role, "manual")
        self.assertEqual(target.category, "dog")
        self.assertEqual(self.w.review.store.get(target.id)["status"], "false_negative")
        self.assertEqual(self.w.review.reason_status, "false_negative")
        self.key(Qt.Key.Key_Return)
        self.assertEqual(len(view.overlay_items["manual"]), 1)
        self.w.annotation_panel.delete_button.click()
        self.assertNotIn(target.id, self.w.review.store.manual_targets)
        self.w.review.undo()
        self.assertIn(target.id, self.w.review.store.manual_targets)
        self.assertTrue(self.w.review.save(explicit=True))
        loaded = AnnotationStore.load(self.folder)
        self.assertEqual(loaded.manual_targets[target.id].category, "dog")

    def test_imported_config_remaps_same_keys_to_status_then_reason(self):
        config_file = self.folder / "reason_config.json"
        config_file.write_text(json.dumps({"reasons": {"false_positive": ["custom reason"]}, "auto_advance": False,
            "shortcuts": {"decision_1": "Q", "decision_2": "W", "decision_3": "E"}}), encoding="utf-8")
        with patch("detretlabel.review_controller.QFileDialog.getOpenFileName", return_value=(str(config_file), "JSON")):
            self.w.annotation_panel.config_button.click()
        self.key(Qt.Key.Key_W)
        self.assertEqual(self.w.review.reason_status, "false_positive")
        self.assertIn("Q", self.w.annotation_panel.wheel.buttons["true_positive"].text())
        self.assertIn("W", self.w.annotation_panel.wheel.buttons["false_positive"].text())
        self.assertNotIn("false_negative", self.w.annotation_panel.wheel.buttons)
        self.key(Qt.Key.Key_Q)
        self.assertEqual(self.w.review.store.get(self.w.review.current.id)["reason"], "custom reason")
        self.assertIn("Q", self.w.annotation_panel.wheel.buttons["true_positive"].text())

    def test_save_failure_keeps_data_and_prevents_close(self):
        self.w.annotation_panel.auto_next.setChecked(False)
        self.key(Qt.Key.Key_1)
        target_id = self.w.review.current.id
        with patch("detretlabel.annotation.os.replace", side_effect=PermissionError("test")), \
             patch("detretlabel.review_controller.QMessageBox.warning"):
            self.assertFalse(self.w.review.save(explicit=True))
            self.w.close()
            self.assertTrue(self.w.isVisible())
            self.assertTrue(self.w.review.store.dirty)
        self.assertTrue(self.w.review.store.completed(target_id))
        self.assertTrue(self.w.review.save(explicit=True))

    def test_reason_mode_escape_keeps_selected_status_but_returns_to_status_keys(self):
        self.w.annotation_panel.auto_next.setChecked(False)
        self.key(Qt.Key.Key_2)
        self.key(Qt.Key.Key_Escape)
        self.assertIsNone(self.w.review.reason_status)
        self.assertEqual(self.w.review.store.get(self.w.review.current.id)["status"], "false_positive")
        self.key(Qt.Key.Key_1)
        self.assertEqual(self.w.review.store.get(self.w.review.current.id)["status"], "true_positive")

    def test_custom_reason_enter_returns_focus_to_review_keys(self):
        self.key(Qt.Key.Key_2)
        first = self.w.review.current.id
        editor = self.w.annotation_panel.reason_combo.lineEdit()
        editor.setFocus()
        QTest.keyClicks(editor, "custom")
        QTest.keyClick(editor, Qt.Key.Key_Return)
        self.assertEqual(self.w.review.store.get(first)["reason"], "custom")
        second = self.w.review.current.id
        self.assertNotEqual(first, second)
        self.assertIs(QApplication.focusWidget(), self.w.image_view)
        QTest.keyClick(self.w.image_view, Qt.Key.Key_1)
        self.assertEqual(self.w.review.store.get(second)["status"], "true_positive")

    def test_drawing_mode_blocks_classification_until_cancelled(self):
        target_id = self.w.review.current.id
        self.w.review.all_positive()
        self.key(Qt.Key.Key_B)
        self.assertTrue(self.w.image_view.draw_mode)
        self.key(Qt.Key.Key_2)
        self.assertEqual(self.w.review.store.get(target_id)["status"], "true_positive")
        self.key(Qt.Key.Key_Escape)
        self.assertFalse(self.w.image_view.draw_mode)
        self.assertEqual(self.w.review.phase, "manual")

    def test_reason_layout_keeps_wheel_and_buttons_readable(self):
        self.key(Qt.Key.Key_2)
        QTest.qWait(30)
        self.app.processEvents()
        panel = self.w.annotation_panel
        self.assertGreater(panel.reason_frame.geometry().top(), panel.wheel.geometry().bottom())
        self.assertTrue(all(button.height() >= 28 for button in panel.reason_buttons))

    def test_status_button_numbers_remain_visible_during_mouse_and_keyboard_review(self):
        panel, review = self.w.annotation_panel, self.w.review
        panel.auto_next.setChecked(False)
        original = {status: button.text() for status, button in panel.wheel.buttons.items()}
        for index, text in enumerate(original.values(), 1):
            self.assertIn(str(index), text)
        for status in ("false_positive", "true_positive"):
            panel.wheel.buttons[status].click()
            self.assertEqual({key: button.text() for key, button in panel.wheel.buttons.items()}, original)
            self.assertTrue(panel.wheel.buttons[status].isChecked())
        self.key(Qt.Key.Key_2)
        self.assertIn("1–9 选原因", panel.stage.text())
        self.key(Qt.Key.Key_1)
        self.assertEqual(review.store.get(review.current.id)["reason"], "种类检错")
        self.assertEqual({key: button.text() for key, button in panel.wheel.buttons.items()}, original)

    def test_reason_expansion_and_long_results_keep_upper_controls_stationary(self):
        panel, review = self.w.annotation_panel, self.w.review
        panel.auto_next.setChecked(False)
        def settle():
            self.app.processEvents()
            QTest.qWait(30)
            self.app.processEvents()
        controls = (panel.scope_combo, panel.show_all_targets, panel.auto_next, panel.previous_button,
                    panel.next_button, panel.unreviewed_button, panel.delete_button,
                    panel.clear_button, panel.save_state, panel.help_label, panel.summary_frame,
                    panel.current_result, panel.note, panel.wheel, panel.gt_button, *panel.wheel.buttons.values())
        def positions():
            return [(widget.mapTo(panel.viewport(), QPoint()), widget.size()) for widget in controls]
        for width, height in ((1440, 960), (1040, 720)):
            review.reason_status = None
            review.render()
            self.w.resize(width, height)
            settle()
            panel.verticalScrollBar().setValue(max(0, panel.wheel.y() - 160))
            expected = positions()
            self.assertLess(panel.summary_frame.geometry().bottom(), panel.wheel.y())
            panel.wheel.buttons["false_positive"].click()
            settle()
            self.assertTrue(panel.reason_frame.isVisible())
            self.assertEqual(positions(), expected)
            panel.wheel.buttons["true_positive"].click()
            settle()
            self.assertEqual(positions(), expected)
            panel.wheel.buttons["false_positive"].click()
            long_reason = "多行原因\n" + "很长的自定义原因" * 30
            review.choose_reason(long_reason)
            review.save(explicit=True)
            settle()
            self.assertIn(long_reason, panel.current_result.toolTip())
            self.assertNotIn("\n", panel.current_result.text())
            self.assertEqual(positions(), expected)
            self.assertFalse(panel.reason_frame.isVisible())
            panel.wheel.buttons["false_positive"].click()
            settle()
            self.key(Qt.Key.Key_Escape)
            settle()
            self.assertEqual(positions(), expected)

    def test_decode_failure_prevents_accidental_classification(self):
        (self.folder / "01.jpg").write_bytes(b"not a JPEG")
        self.w.scan_folder(self.folder)
        self.wait()
        target_id = self.w.review.current.id
        self.assertIsNone(self.w.image_view.path)
        self.key(Qt.Key.Key_1)
        self.assertFalse(self.w.review.store.completed(target_id))

    def test_default_overlays_follow_only_current_target_and_gt_scope(self):
        gt = [{"name": "01.jpg", "objects": [{"xywh": [12, 22, 100, 150], "type": "GT person"}]}]
        self.import_source(gt, "ground_truth")
        panel, view = self.w.annotation_panel, self.w.image_view
        self.assertFalse(panel.show_all_targets.isChecked())
        self.assertTrue(panel.keep_reviewed_targets.isChecked())
        self.assertFalse(panel.show_noncurrent_labels.isChecked())
        self.assertTrue(panel.show_noncurrent_labels.isEnabled())
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [True, False])
        self.assertEqual([item.isVisible() for item in view.label_items["algorithm"]], [True, False])
        self.assertFalse(view.overlay_items["ground_truth"][0].isVisible())
        self.key(Qt.Key.Key_Down)
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [False, True])
        self.assertEqual([item.isVisible() for item in view.label_items["algorithm"]], [False, True])
        self.w.review.all_positive()
        panel.scope_combo.setCurrentIndex(1)
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["algorithm"]))
        self.assertFalse(view.overlay_items["ground_truth"][0].isVisible())
        self.key(Qt.Key.Key_G)
        self.assertTrue(view.label_items["ground_truth"][0].isVisible())
        self.w.clear_json("ground_truth")
        self.assertIsNone(self.w.review.current)
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["algorithm"]))

    def test_all_algorithm_target_and_noncurrent_label_options_are_dependent(self):
        gt = [{"name": "01.jpg", "objects": [{"xywh": [12, 22, 100, 150], "type": "GT person"}]}]
        self.import_source(gt, "ground_truth")
        panel, view = self.w.annotation_panel, self.w.image_view
        view.auto_fit = False
        view.scale(1.5, 1.5)
        zoom = view.transform()
        panel.show_all_targets.click()
        self.assertTrue(panel.show_noncurrent_labels.isEnabled())
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["algorithm"]))
        self.assertFalse(view.overlay_items["ground_truth"][0].isVisible())
        self.assertEqual([item.isVisible() for item in view.label_items["algorithm"]], [True, False])
        self.assertFalse(view.label_items["ground_truth"][0].isVisible())
        view._layout_overlay_labels()
        self.assertFalse(view.label_items["ground_truth"][0].isVisible())
        panel.show_noncurrent_labels.click()
        self.assertTrue(all(item.isVisible() for item in view.label_items["algorithm"]))
        self.assertFalse(view.label_items["ground_truth"][0].isVisible())
        panel.show_all_targets.click()
        self.assertTrue(panel.show_noncurrent_labels.isEnabled())
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [True, False])
        self.assertEqual([item.isVisible() for item in view.label_items["algorithm"]], [True, False])
        self.assertFalse(view.label_items["ground_truth"][0].isVisible())
        self.assertEqual(view.transform(), zoom)
        panel.show_all_targets.click()
        self.assertTrue(all(item.isVisible() for item in view.label_items["algorithm"]))
        self.assertFalse(view.label_items["ground_truth"][0].isVisible())

    def test_noncurrent_label_visibility_ignores_review_status(self):
        panel, review, view = self.w.annotation_panel, self.w.review, self.w.image_view
        panel.keep_reviewed_targets.setChecked(False)
        panel.show_all_targets.setChecked(True)
        first = review.current.id
        self.key(Qt.Key.Key_1)
        self.assertNotEqual(review.current.id, first)
        self.assertEqual([item.isVisible() for item in view.label_items["algorithm"]], [False, True])
        self.key(Qt.Key.Key_Up)
        self.key(Qt.Key.Key_Backspace)
        self.key(Qt.Key.Key_Down)
        self.assertEqual([item.isVisible() for item in view.label_items["algorithm"]], [False, True])
        self.key(Qt.Key.Key_Z, Qt.KeyboardModifier.ControlModifier)
        self.key(Qt.Key.Key_Down)
        self.assertEqual([item.isVisible() for item in view.label_items["algorithm"]], [False, True])
        self.key(Qt.Key.Key_Up)
        self.key(Qt.Key.Key_Backspace)
        panel.note.setFocus()
        QTest.keyClicks(panel.note, "note only")
        review.flush_note()
        self.key(Qt.Key.Key_Down)
        self.assertEqual([item.isVisible() for item in view.label_items["algorithm"]], [False, True])
        panel.show_noncurrent_labels.setChecked(True)
        self.assertEqual([item.isVisible() for item in view.label_items["algorithm"]], [True, True])

    def test_keep_reviewed_overlays_are_independent_of_all_and_noncurrent_switches(self):
        data = copy.deepcopy(self.data)
        data["results"][0]["detections"].append(
            {"bbox": [360, 60, 430, 180], "type": "dog", "conf": .7})
        self.import_source(data)
        panel, review, view = self.w.annotation_panel, self.w.review, self.w.image_view
        view.auto_fit = False
        view.scale(1.2, 1.2)
        zoom = view.transform()
        self.key(Qt.Key.Key_1)
        self.assertTrue(panel.keep_reviewed_targets.isChecked())
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [True, True, False])
        panel.keep_reviewed_targets.click()
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [False, True, False])
        panel.keep_reviewed_targets.click()
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [True, True, False])
        self.assertEqual([item.isVisible() for item in view.label_items["algorithm"]], [True, True, False])
        QTest.keyPress(view, Qt.Key.Key_H)
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["algorithm"]))
        self.assertEqual([item.isVisible() for item in view.label_items["algorithm"]], [True, True, False])
        QTest.keyRelease(view, Qt.Key.Key_H)
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [True, True, False])
        panel.show_all_targets.setChecked(True)
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["algorithm"]))
        self.assertEqual([item.isVisible() for item in view.label_items["algorithm"]], [True, True, False])
        panel.show_noncurrent_labels.setChecked(True)
        self.assertTrue(all(item.isVisible() for item in view.label_items["algorithm"]))
        panel.show_all_targets.setChecked(False)
        panel.show_noncurrent_labels.setChecked(False)
        review.clicked_target("algorithm", 3)
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [True, False, True])
        self.assertEqual([item.isVisible() for item in view.label_items["algorithm"]], [True, False, True])
        panel.keep_reviewed_targets.click()
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [False, False, True])
        self.assertEqual([item.isVisible() for item in view.label_items["algorithm"]], [False, False, True])
        self.assertEqual(view.transform(), zoom)

    def test_current_target_and_label_use_manual_purple_and_restore_source_color_on_navigation(self):
        panel, review, view = self.w.annotation_panel, self.w.review, self.w.image_view
        self.matching_gt()
        self.wait_matches()
        panel.show_all_targets.setChecked(True)
        panel.show_noncurrent_labels.setChecked(True)
        view.auto_fit = False
        view.scale(1.3, 1.3)
        zoom = view.transform()
        source_style = BoxStyle("#ff00aa", "dot")
        with patch("detretlabel.app.BoxStyleDialog") as dialog:
            dialog.return_value.exec.return_value = QDialog.DialogCode.Accepted
            dialog.return_value.selected_style.return_value = source_style
            self.w.configure_box_style("algorithm")
        self.assertEqual([item.pen().color().name() for item in view.overlay_items["algorithm"]],
                         [MANUAL_BOX_COLOR, source_style.color])
        self.assertEqual([label.color.name() for label in view.label_items["algorithm"]],
                         [MANUAL_BOX_COLOR, source_style.color])
        self.key(Qt.Key.Key_1)
        self.assertEqual(review.current.category, "car")
        self.wait_matches()
        self.assertEqual([item.pen().color().name() for item in view.overlay_items["algorithm"]],
                         [source_style.color, MANUAL_BOX_COLOR])
        self.assertEqual([label.color.name() for label in view.label_items["algorithm"]],
                         [source_style.color, MANUAL_BOX_COLOR])
        self.assertTrue(all(item.pen().style() == Qt.PenStyle.DotLine for item in view.overlay_items["algorithm"]))
        self.w.target_tables["ground_truth"].setCurrentCell(3, 0)
        self.assertEqual(review.current.category, "car")
        self.assertEqual(view.overlay_items["algorithm"][1].pen().color().name(), MANUAL_BOX_COLOR)
        self.key(Qt.Key.Key_Up)
        self.assertEqual([label.color.name() for label in view.label_items["algorithm"]],
                         [MANUAL_BOX_COLOR, source_style.color])
        self.assertEqual(self.w.box_styles["algorithm"], source_style)
        self.assertEqual(view.transform(), zoom)

    def test_target_review_updates_only_its_image_progress_row(self):
        table = self.w.image_table
        changes = QSignalSpy(table.model().dataChanged)
        self.key(Qt.Key.Key_1)
        self.assertGreater(changes.count(), 0)
        self.assertEqual({changes.at(i)[0].row() for i in range(changes.count())}, {0})
        self.assertTrue(table.item(0, 3).text().startswith("1/2"))
        self.w.move_image(1)
        changes = QSignalSpy(table.model().dataChanged)
        self.key(Qt.Key.Key_1)
        self.assertGreater(changes.count(), 0)
        self.assertEqual({changes.at(i)[0].row() for i in range(changes.count())}, {1})
        self.assertTrue(table.item(0, 3).text().startswith("1/2"))
        self.assertTrue(table.item(1, 3).text().startswith("1/1"))

    def test_retained_targets_require_finished_review_and_follow_clear_and_undo(self):
        panel, review, view = self.w.annotation_panel, self.w.review, self.w.image_view
        panel.keep_reviewed_targets.setChecked(True)
        first, second = review.queue()
        self.key(Qt.Key.Key_2)
        self.key(Qt.Key.Key_Down)
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [False, True])
        panel.note.setFocus()
        QTest.keyClicks(panel.note, "note only")
        review.flush_note()
        review.select(first.id)
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [True, False])
        panel.confirm_button.click()
        self.assertTrue(review.target_completed(first))
        self.assertEqual(review.current.id, second.id)
        self.assertTrue(all(item.isVisible() for item in view.label_items["algorithm"]))
        review.select(first.id)
        panel.clear_button.click()
        review.select(second.id)
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [False, True])
        panel.undo_button.click()
        review.select(second.id)
        self.assertTrue(all(item.isVisible() for item in view.label_items["algorithm"]))
        with patch("detretlabel.review_controller.QMessageBox.question", return_value=QMessageBox.StandardButton.Yes):
            panel.clear_image_button.click()
        self.assertTrue(panel.keep_reviewed_targets.isChecked())
        self.assertEqual(sum(item.isVisible() for item in view.overlay_items["algorithm"]), 1)
        self.assertEqual(sum(item.isVisible() for item in view.label_items["algorithm"]), 1)

    def test_retained_overlays_use_saved_reviews_and_sorted_rows_for_each_image(self):
        data = copy.deepcopy(self.data)
        data["results"][0]["detections"][0]["conf"] = .2
        self.import_source(data)
        panel, review = self.w.annotation_panel, self.w.review
        panel.keep_reviewed_targets.setChecked(True)
        review.select(review.queue()[1].id)
        self.key(Qt.Key.Key_1)
        self.assertEqual(review.current.category, "car")
        self.assertTrue(review.save(explicit=True))
        self.w.move_image(1)
        self.assertEqual(self.w.image_view.retained_algorithm_indices, set())
        self.w.move_image(-1)
        self.assertEqual(self.w.image_view.retained_algorithm_indices, {2})
        self.assertTrue(all(item.isVisible() for item in self.w.image_view.label_items["algorithm"]))
        self.w.close()
        self.w = MainWindow(self.settings)
        self.w.show()
        self.w.scan_folder(self.folder)
        self.wait()
        self.import_source(data)
        self.w.annotation_panel.keep_reviewed_targets.setChecked(True)
        self.assertEqual(self.w.review.current.category, "car")
        self.assertEqual(self.w.image_view.retained_algorithm_indices, {2})
        self.assertTrue(all(item.isVisible() for item in self.w.image_view.label_items["algorithm"]))

    def test_retained_labels_honor_source_style_and_survive_reference_and_review_updates(self):
        self.matching_gt()
        self.wait_matches()
        panel, review, view = self.w.annotation_panel, self.w.review, self.w.image_view
        panel.keep_reviewed_targets.setChecked(True)
        self.key(Qt.Key.Key_1)
        self.wait_matches()
        review.gt_reference.inspect(1, "ground_truth")
        self.assertTrue(all(item.isVisible() for item in view.label_items["algorithm"]))
        self.assertEqual([item.isVisible() for item in view.overlay_items["ground_truth"]], [True, False, False, False])
        style = replace(self.w.box_styles["algorithm"], show_category=False, show_score=False)
        with patch("detretlabel.app.BoxStyleDialog") as dialog:
            dialog.return_value.exec.return_value = QDialog.DialogCode.Accepted
            dialog.return_value.selected_style.return_value = style
            self.w.source_panels["algorithm"].box_style.click()
        self.assertFalse(view.label_items["algorithm"])
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["algorithm"]))
        review.all_positive()
        self.assertEqual(review.phase, "manual")
        panel.show_all_targets.setChecked(False)
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["algorithm"]))
        self.assertFalse(panel.gt_button.isChecked())
        self.assertFalse(any(item.isVisible() for item in view.overlay_items["ground_truth"]))
        panel.keep_reviewed_targets.setChecked(False)
        self.assertFalse(any(item.isVisible() for item in view.overlay_items["algorithm"]))

    def test_gt_toggle_is_independent_and_shows_or_hides_all_gt_together(self):
        panel, view, review = self.w.annotation_panel, self.w.image_view, self.w.review
        self.assertFalse(panel.gt_button.isEnabled())
        self.key(Qt.Key.Key_G)
        self.assertFalse(panel.gt_button.isChecked())
        gt = [{"name": "01.jpg", "objects": [
            {"xywh": [12, 22, 100, 150], "type": "GT person"},
            {"xywh": [220, 50, 80, 140], "type": "GT car"}]}]
        self.import_source(gt, "ground_truth")
        self.assertTrue(panel.gt_button.isEnabled())
        self.assertFalse(panel.gt_button.isChecked())
        target_id = review.current.id
        self.key(Qt.Key.Key_G)
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["ground_truth"]))
        self.assertTrue(all(item.isVisible() for item in view.label_items["ground_truth"]))
        for enabled in (True, False, True):
            panel.show_all_targets.setChecked(enabled)
            panel.show_noncurrent_labels.setChecked(enabled)
            self.assertTrue(all(item.isVisible() for item in view.overlay_items["ground_truth"]))
            self.assertTrue(all(item.isVisible() for item in view.label_items["ground_truth"]))
        self.assertEqual(review.current.id, target_id)
        self.assertEqual(panel.scope_combo.currentData(), "algorithm")
        self.key(Qt.Key.Key_Right)
        self.key(Qt.Key.Key_Left)
        self.assertTrue(panel.gt_button.isChecked())
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["ground_truth"]))
        panel.gt_button.click()
        self.assertFalse(any(item.isVisible() for item in view.overlay_items["ground_truth"]))
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["algorithm"]))
        self.assertGreater(panel.gt_button.height(), panel.clear_button.height())
        self.assertGreater(panel.gt_button.geometry().left(), panel.wheel.buttons["true_positive"].geometry().right())
        self.assertLess(panel.gt_button.y(), panel.wheel.center.y())
        self.w.clear_json("ground_truth")
        self.assertFalse(panel.gt_button.isEnabled())

    def test_gt_shortcut_can_be_configured_and_does_not_fire_during_text_input(self):
        self.import_source([{"name": "01.jpg", "bbox": [10, 20, 110, 170], "type": "person"}], "ground_truth")
        config_path = self.folder / "gt_shortcut.json"
        config_path.write_text(json.dumps({"shortcuts": {"toggle_ground_truth": "T"}}), encoding="utf-8")
        with patch("detretlabel.review_controller.QFileDialog.getOpenFileName", return_value=(str(config_path), "JSON")):
            self.w.annotation_panel.config_button.click()
        panel, review = self.w.annotation_panel, self.w.review
        self.assertIn("T", panel.gt_button.text())
        self.key(Qt.Key.Key_G)
        self.assertFalse(panel.gt_button.isChecked())
        self.key(Qt.Key.Key_T)
        self.assertTrue(panel.gt_button.isChecked())
        panel.note.setFocus()
        QTest.keyClicks(panel.note, "t")
        self.assertEqual(panel.note.text(), "t")
        self.assertTrue(panel.gt_button.isChecked())
        self.key(Qt.Key.Key_2)
        target_id = review.current.id
        self.key(Qt.Key.Key_T)
        self.assertFalse(panel.gt_button.isChecked())
        self.assertEqual(review.reason_status, "false_positive")
        self.assertEqual(review.current.id, target_id)

    def test_manual_boxes_remain_visible_independently_of_algorithm_and_gt_switches(self):
        review, panel, view = self.w.review, self.w.annotation_panel, self.w.image_view
        for box, category in (((30, 40, 100, 120), "dog"), ((130, 140, 200, 220), "cat")):
            target = ReviewTarget.manual("01.jpg", box, category)
            review.store.set_annotation(target, "false_negative")
        review.rebuild()
        review.fill_manual_table()
        self.w._update_overlays()
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["manual"]))
        self.assertTrue(all(item.isVisible() for item in view.label_items["manual"]))
        panel.show_all_targets.setChecked(True)
        panel.show_noncurrent_labels.setChecked(True)
        panel.show_all_targets.setChecked(False)
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["manual"]))
        self.assertTrue(all(item.isVisible() for item in view.label_items["manual"]))

    def test_overlay_filters_survive_image_changes_and_honor_source_label_settings(self):
        panel, view = self.w.annotation_panel, self.w.image_view
        panel.show_all_targets.setChecked(True)
        panel.show_noncurrent_labels.setChecked(True)
        style = replace(self.w.box_styles["algorithm"], show_category=False, show_score=False)
        with patch("detretlabel.app.BoxStyleDialog") as dialog:
            dialog.return_value.exec.return_value = QDialog.DialogCode.Accepted
            dialog.return_value.selected_style.return_value = style
            self.w.source_panels["algorithm"].box_style.click()
        self.assertFalse(view.label_items["algorithm"])
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["algorithm"]))
        self.key(Qt.Key.Key_Right)
        self.assertEqual(self.w.current_image().name, "02.png")
        self.assertFalse(view.label_items["algorithm"])
        self.key(Qt.Key.Key_Left)
        self.assertFalse(view.label_items["algorithm"])
        panel.show_all_targets.setChecked(False)
        self.key(Qt.Key.Key_Right)
        self.key(Qt.Key.Key_Left)
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [True, False])

    def test_hold_all_keyboard_mouse_and_label_setting_restore_original_display(self):
        panel, view = self.w.annotation_panel, self.w.image_view
        panel.show_noncurrent_labels.setChecked(True)
        view.auto_fit = False
        view.scale(1.4, 1.4)
        zoom = view.transform()
        QTest.keyPress(view, Qt.Key.Key_H)
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["algorithm"]))
        self.assertTrue(all(item.isVisible() for item in view.label_items["algorithm"]))
        self.assertFalse(panel.show_all_targets.isChecked())
        QTest.keyRelease(view, Qt.Key.Key_H)
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [True, False])
        QTest.mousePress(panel.hold_button, Qt.MouseButton.LeftButton)
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["algorithm"]))
        QTest.mouseRelease(panel.hold_button, Qt.MouseButton.LeftButton)
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [True, False])
        panel.show_noncurrent_labels.setChecked(False)
        QTest.keyPress(view, Qt.Key.Key_H)
        self.assertEqual([item.isVisible() for item in view.label_items["algorithm"]], [True, False])
        QApplication.sendEvent(self.w, QEvent(QEvent.Type.WindowDeactivate))
        self.assertFalse(self.w.review._hold_sources)
        self.assertEqual([item.isVisible() for item in view.overlay_items["algorithm"]], [True, False])
        QTest.keyRelease(view, Qt.Key.Key_H)
        self.assertEqual(view.transform(), zoom)
        self.assertEqual(panel.hold_button.size(), panel.gt_button.size())

    def test_hold_shortcut_is_configurable_and_does_not_fire_while_typing(self):
        config = ReviewConfig()
        config.shortcuts["hold_all_algorithm"] = "Ctrl+H"
        self.w.review.apply_config(config)
        view, panel = self.w.image_view, self.w.annotation_panel
        panel.note.setFocus()
        QTest.keyClicks(panel.note, "h")
        self.assertFalse(self.w.review._hold_sources)
        view.setFocus()
        QTest.keyPress(view, Qt.Key.Key_H, Qt.KeyboardModifier.ControlModifier)
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["algorithm"]))
        QTest.keyRelease(view, Qt.Key.Key_Control)
        self.assertFalse(self.w.review._hold_sources)
        QTest.keyRelease(view, Qt.Key.Key_H)

    def test_all_positive_updates_only_current_algorithm_targets_and_undoes_once(self):
        review, panel = self.w.review, self.w.annotation_panel
        panel.auto_next.setChecked(False)
        first, second = review.queue()
        review.choose_status("false_positive")
        review.choose_reason("wrong class")
        panel.note.setFocus()
        QTest.keyClicks(panel.note, "keep note")
        review.flush_note()
        manual = ReviewTarget.manual("01.jpg", (30, 40, 100, 130), "dog")
        review.store.set_annotation(manual, "false_negative", "missed")
        review.rebuild()
        history_size = len(review.store.history)
        self.key(Qt.Key.Key_1, Qt.KeyboardModifier.ControlModifier)
        self.assertEqual(review.phase, "manual")
        self.assertEqual(len(review.store.history), history_size + 1)
        self.assertEqual(review.store.get(first.id)["status"], "true_positive")
        self.assertEqual(review.store.get(first.id)["reason"], "")
        self.assertEqual(review.store.get(first.id)["note"], "keep note")
        self.assertEqual(review.store.get(second.id)["status"], "true_positive")
        self.assertEqual(review.store.get(manual.id)["status"], "false_negative")
        self.assertFalse(review.store.get(review.role_targets("algorithm", self.folder / "02.png")[0].id))
        self.key(Qt.Key.Key_Z, Qt.KeyboardModifier.ControlModifier)
        self.assertEqual(review.store.get(first.id)["status"], "false_positive")
        self.assertEqual(review.store.get(first.id)["reason"], "wrong class")
        self.assertIsNone(review.store.get(second.id))

    def test_review_can_start_with_unfinished_algorithm_and_preserves_pending_reason(self):
        review, panel = self.w.review, self.w.annotation_panel
        self.assertTrue(review.start_review())
        self.assertFalse(review.algorithm_ready())
        review.leave_review()
        self.key(Qt.Key.Key_1)
        self.key(Qt.Key.Key_2)
        self.assertFalse(review.algorithm_ready())
        panel.reason_combo.setCurrentText("先复核，稍后确认原因")
        target_id = review.current.id
        self.assertTrue(review.start_review())
        self.assertEqual(review.phase, "manual")
        self.assertEqual(review.store.get(target_id)["reason_draft"], "先复核，稍后确认原因")
        review.leave_review()
        self.assertEqual(review.current.id, target_id)
        self.assertEqual(panel.reason_combo.currentText(), "先复核，稍后确认原因")
        self.key(Qt.Key.Key_Return)
        self.assertTrue(panel.review_button.isEnabled())
        self.assertEqual(review.phase, "manual")
        self.assertFalse(any(button.isEnabled() for button in panel.wheel.buttons.values()))
        self.assertIsNone(review.current)

    def test_draw_is_only_available_in_stage_four_and_leaving_cancels_drawing(self):
        review, panel, view = self.w.review, self.w.annotation_panel, self.w.image_view
        self.assertFalse(hasattr(panel, "draw_button"))
        self.assertFalse(panel.review_draw_button.isVisible())
        self.key(Qt.Key.Key_B)
        self.assertEqual(review.phase, "algorithm")
        self.assertFalse(view.draw_mode)
        with patch("detretlabel.review_controller.QInputDialog.getText") as dialog:
            review.add_manual((1, 2, 30, 40))
        dialog.assert_not_called()
        self.assertTrue(panel.review_button.isEnabled())
        self.key(Qt.Key.Key_R)
        self.assertEqual(review.phase, "manual")
        self.assertFalse(review.algorithm_ready())
        self.assertTrue(panel.review_draw_button.isVisible())
        self.key(Qt.Key.Key_B)
        self.assertTrue(view.draw_mode)
        panel.scope_combo.setCurrentIndex(0)
        self.assertEqual(review.phase, "algorithm")
        self.assertFalse(view.draw_mode)
        self.assertFalse(panel.review_draw_button.isVisible())

    def test_partial_algorithm_can_finish_miss_review_save_and_resume_then_require_recheck(self):
        review = self.w.review
        first_id = review.current.id
        self.key(Qt.Key.Key_1)
        self.assertFalse(review.algorithm_ready())
        self.assertTrue(review.start_review())
        with patch("detretlabel.review_controller.QInputDialog.getText", return_value=("人工漏检", True)):
            review.add_manual((300, 10, 350, 80))
        manual_id = review.current.id
        review.choose_reason("完全未检出")
        self.assertTrue(self.w.annotation_panel.finish_review_button.isEnabled())
        review.finish_review()
        self.assertEqual(self.w.current_image().name, "02.png")
        image = self.folder / "01.jpg"
        self.assertTrue(review.miss_review_completed(image))
        self.assertFalse(review.review_completed(image))
        self.assertIn("03未完成", self.w.image_table.item(0, 3).text())
        self.assertIn("04已完成", self.w.image_table.item(0, 3).text())
        report = review.report_snapshot()
        self.assertEqual(report.images[0]["algorithm_state"], "未完成")
        self.assertEqual(report.images[0]["miss_review_state"], "已完成")
        self.w.close()
        self.w = MainWindow(self.settings)
        self.w.show()
        self.w.scan_folder(self.folder)
        self.wait()
        self.import_source(self.data)
        resumed = self.w.review
        self.assertEqual(resumed.store.get(first_id)["status"], "true_positive")
        self.assertEqual(resumed.store.get(manual_id)["reason"], "完全未检出")
        self.assertTrue(resumed.miss_review_completed(image))
        self.assertFalse(resumed.review_completed(image))
        self.assertEqual(resumed.current.category, "car")
        self.key(Qt.Key.Key_1)
        self.assertTrue(resumed.algorithm_ready())
        self.assertFalse(resumed.miss_review_completed())
        self.assertEqual(resumed.phase, "manual")
        self.assertIn(manual_id, resumed.store.manual_targets)

    def test_partial_no_miss_confirmation_keeps_image_unfinished_and_returns_to_stage_three(self):
        review = self.w.review
        self.w.search.setText("01")
        review.start_review()
        self.assertIn("确认无漏检", self.w.annotation_panel.finish_review_button.text())
        review.finish_review()
        self.assertEqual(self.w.current_image().name, "01.jpg")
        self.assertEqual(review.phase, "algorithm")
        self.assertFalse(review.algorithm_ready())
        self.assertTrue(review.miss_review_completed())
        self.assertFalse(review.review_completed())
        self.assertFalse(review.store.records)
        self.assertTrue(AnnotationStore.load(self.folder).review_completed("01.jpg", review.algorithm_signature()))
        self.assertIn("03未完成", self.w.image_table.item(0, 3).text())
        self.assertIn("03 尚未完成", self.w.statusBar().currentMessage())

    def test_pending_reason_draft_and_note_survive_switch_image_and_restart(self):
        review, panel = self.w.review, self.w.annotation_panel
        review.choose_status("false_positive")
        target_id = review.current.id
        panel.reason_combo.setCurrentText("待核实的自定义原因")
        panel.note.setText("未完成图片备注")
        panel.note_edited.emit(panel.note.text())
        self.w.image_table.setCurrentCell(1, 0)
        self.assertEqual(review.store.get(target_id)["reason_draft"], "待核实的自定义原因")
        self.w.close()
        self.w = MainWindow(self.settings)
        self.w.show()
        self.w.scan_folder(self.folder)
        self.wait()
        self.import_source(self.data)
        review, panel = self.w.review, self.w.annotation_panel
        self.assertEqual(review.current.id, target_id)
        self.assertEqual(review.reason_status, "false_positive")
        self.assertEqual(panel.reason_combo.currentText(), "待核实的自定义原因")
        self.assertEqual(panel.note.text(), "未完成图片备注")
        self.assertFalse(review.target_completed(review.current))
        self.assertFalse(review.algorithm_ready())
        self.assertIn("原因待确认", self.w.target_table.item(0, 7).text())
        self.assertTrue(review.start_review())
        review.finish_review()
        snapshot = review.report_snapshot()
        detail = next(row for row in snapshot.targets if row["target_id"] == target_id)
        self.assertEqual(detail["draft"], "待核实的自定义原因")
        self.assertEqual(detail["editing"], "原因待确认")
        self.assertEqual(snapshot.images[0]["pending_reasons"], 1)

    def test_review_never_forces_gt_and_gt_changes_do_not_invalidate_completion(self):
        review, panel, view = self.w.review, self.w.annotation_panel, self.w.image_view
        self.import_source([{"name": "01.jpg", "bbox": [20, 20, 100, 150], "type": "gt"}], "ground_truth")
        review.all_positive()
        view.auto_fit = False
        view.scale(1.3, 1.3)
        zoom = view.transform()
        self.assertTrue(review.start_review())
        self.assertFalse(panel.gt_button.isChecked())
        self.assertFalse(any(item.isVisible() for item in view.overlay_items["ground_truth"]))
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["algorithm"]))
        self.assertEqual(view.transform(), zoom)
        self.key(Qt.Key.Key_G)
        self.assertTrue(all(item.isVisible() for item in view.overlay_items["ground_truth"]))
        self.key(Qt.Key.Key_G)
        review.finish_review()
        self.assertTrue(review.review_completed(self.folder / "01.jpg"))
        self.import_source([{"name": "01.jpg", "bbox": [120, 20, 200, 150], "type": "different"}], "ground_truth")
        self.assertTrue(review.review_completed(self.folder / "01.jpg"))
        self.assertFalse(panel.gt_button.isChecked())
        self.assertFalse(any(t.role == "ground_truth" for t in review.queue()))
        review.leave_review()
        self.assertFalse(panel.show_all_targets.isChecked())

    def test_zero_detection_image_needs_explicit_review_and_restores_after_reopen(self):
        self.w.image_table.setCurrentCell(2, 0)
        review = self.w.review
        self.assertFalse(review.review_completed())
        self.assertIn("0/0", self.w.image_table.item(2, 3).text())
        self.assertTrue(review.start_review())
        self.assertFalse(review.store.records)
        review.finish_review()
        self.assertTrue(review.review_completed(self.folder / "03.jpg"))
        self.assertIn("复核完成", self.w.image_table.item(2, 3).text())
        self.w.close()
        self.w = MainWindow(self.settings)
        self.w.show()
        self.w.scan_folder(self.folder)
        self.wait()
        self.import_source(self.data)
        self.w.image_table.setCurrentCell(2, 0)
        self.assertTrue(self.w.review.review_completed())
        self.key(Qt.Key.Key_Z, Qt.KeyboardModifier.ControlModifier)
        self.assertTrue(self.w.review.review_completed())

    def test_manual_miss_review_complete_edit_and_undo_restore_completion(self):
        self.w.image_table.setCurrentCell(2, 0)
        review = self.w.review
        review.start_review()
        with patch("detretlabel.review_controller.QInputDialog.getText", return_value=("dog", True)):
            review.add_manual((30, 40, 100, 120))
        target = review.current
        self.assertEqual(target.role, "manual")
        self.assertTrue(self.w.annotation_panel.reason_frame.parentWidget() is self.w.annotation_panel.review_frame)
        review.choose_reason("too small")
        review.finish_review()
        self.assertTrue(review.review_completed(self.folder / "03.jpg"))
        self.w.image_table.setCurrentCell(2, 0)
        review.start_review()
        review.select(target.id)
        review.choose_reason("occluded")
        self.assertFalse(review.review_completed())
        review.undo()
        self.assertFalse(review.review_completed())
        self.assertEqual(review.store.get(target.id)["reason"], "too small")

    def test_algorithm_change_invalidates_review_and_next_review_finds_zero_target_images(self):
        review = self.w.review
        review.all_positive()
        review.start_review()
        review.finish_review()
        self.w.image_table.setCurrentCell(0, 0)
        review.leave_review()
        first = review.current.id
        review.choose_status("false_positive")
        self.assertFalse(review.review_completed())
        review.undo()
        self.assertTrue(review.review_completed())
        self.assertEqual(review.store.get(first)["status"], "true_positive")
        self.w.image_table.setCurrentCell(1, 0)
        review.all_positive()
        review.start_review()
        review.finish_review()
        self.assertEqual(self.w.current_image().name, "03.jpg")
        self.assertEqual(review.phase, "manual")
        self.assertFalse(review.review_completed())

    def test_finish_shortcut_commits_custom_reason_while_editing_and_accepts_keypad_enter(self):
        config = ReviewConfig()
        config.shortcuts["finish_review"] = "Ctrl+Return"
        self.w.review.apply_config(config)
        self.w.image_table.setCurrentCell(2, 0)
        review = self.w.review
        review.start_review()
        with patch("detretlabel.review_controller.QInputDialog.getText", return_value=("dog", True)):
            review.add_manual((30, 40, 100, 120))
        target_id = review.current.id
        editor = self.w.annotation_panel.reason_combo.lineEdit()
        editor.setFocus()
        QTest.keyClicks(editor, "custom missing")
        QTest.keyClick(editor, Qt.Key.Key_Enter, Qt.KeyboardModifier.ControlModifier)
        self.assertEqual(review.store.get(target_id)["reason"], "custom missing")
        self.assertTrue(review.review_completed(self.folder / "03.jpg"))
        self.assertEqual(self.w.current_image().name, "01.jpg")
        self.assertIs(QApplication.focusWidget(), self.w.image_view)

    def test_default_space_finishes_review_but_remains_text_while_editing(self):
        review, panel = self.w.review, self.w.annotation_panel
        self.assertEqual(review.config.shortcuts["finish_review"], "Space")
        review.all_positive()
        with patch("detretlabel.review_controller.QInputDialog.getText", return_value=("dog", True)):
            review.add_manual((30, 40, 100, 120))
        review.choose_reason("完全未检出")
        panel.note.setFocus()
        QTest.keyClick(panel.note, Qt.Key.Key_Space)
        self.assertEqual(panel.note.text(), " ")
        self.assertFalse(review.review_completed())
        self.assertEqual(self.w.current_image().name, "01.jpg")
        self.key(Qt.Key.Key_Space)
        self.assertTrue(review.review_completed(self.folder / "01.jpg"))
        self.assertEqual(self.w.current_image().name, "02.png")

    def test_completion_undo_and_algorithm_source_change_require_new_review(self):
        review = self.w.review
        review.all_positive()
        review.start_review()
        review.finish_review()
        self.assertTrue(review.review_completed(self.folder / "01.jpg"))
        review.undo()
        self.assertFalse(review.review_completed())
        review.start_review()
        review.finish_review()
        self.assertTrue(review.review_completed(self.folder / "01.jpg"))
        changed = copy.deepcopy(self.data)
        changed["results"][0]["detections"][0]["conf"] = .45
        self.import_source(changed)
        self.assertEqual(review.phase, "algorithm")
        self.assertFalse(review.review_completed())

    def test_last_positive_enters_review_even_when_target_auto_next_is_disabled(self):
        review = self.w.review
        self.w.annotation_panel.auto_next.setChecked(False)
        review.choose_status("true_positive")
        self.assertEqual(review.phase, "algorithm")
        review.move_target(1)
        review.choose_status("true_positive")
        self.assertEqual(review.phase, "manual")
        self.assertIsNone(review.current)
        self.assertFalse(self.w.annotation_panel.gt_button.isChecked())

    def test_last_false_positive_reason_button_enters_review_after_confirmation(self):
        review = self.w.review
        review.choose_status("true_positive")
        review.choose_status("false_positive")
        self.assertEqual(review.phase, "algorithm")
        self.w.annotation_panel.reason_buttons[0].click()
        self.assertEqual(review.phase, "manual")
        self.assertIsNone(review.current)

    def test_finishing_review_saves_then_advances_to_next_algorithm_image(self):
        review, panel = self.w.review, self.w.annotation_panel
        panel.all_positive_button.click()
        self.assertEqual(review.phase, "manual")
        panel.finish_review_button.click()
        self.assertEqual(self.w.current_image().name, "02.png")
        self.assertEqual(review.phase, "algorithm")
        self.assertEqual(review.current.category, "cat")
        saved = json.loads(review.store.path.read_text(encoding="utf-8"))
        self.assertTrue(saved["image_reviews"]["01.jpg"]["completed"])

    def test_finishing_skips_completed_images_and_stays_when_all_are_done(self):
        review = self.w.review
        for target in review.role_targets("algorithm", self.folder / "02.png"):
            review.store.set_annotation(target, "true_positive")
        review.store.set_image_review("02.png", review.algorithm_signature(), True)
        review.all_positive()
        review.finish_review()
        self.assertEqual(self.w.current_image().name, "03.jpg")
        self.assertEqual(review.phase, "manual")
        review.finish_review()
        self.assertEqual(self.w.current_image().name, "03.jpg")
        self.assertTrue(review.review_completed())
        self.assertIn("所有可见图片已完成", self.w.statusBar().currentMessage())

    def test_type_mapping_import_updates_labels_table_and_card_without_reset(self):
        review = self.w.review
        self.w.annotation_panel.auto_next.setChecked(False)
        review.choose_status("false_positive")
        self.w.annotation_panel.reason_combo.setCurrentText("尚未确认的原因")
        target = review.current
        signature = review.algorithm_signature()
        self.w.image_view.scale(1.5, 1.5)
        zoom = self.w.image_view.transform()
        mapping = self.folder / "types.txt"
        mapping.write_text("person:行人\ncar:汽车", encoding="utf-8")
        with patch("detretlabel.app.QFileDialog.getOpenFileName", return_value=(str(mapping), "")):
            self.w.source_panels["algorithm"].type_mapping_button.click()
        self.assertEqual(self.w.image_view.label_items["algorithm"][0].full_text, "行人 · 0.9")
        self.assertEqual(self.w.target_table.item(0, 1).text(), "行人")
        self.assertIn("person", self.w.target_table.item(0, 1).toolTip())
        self.assertIn("行人", self.w.annotation_panel.target_info.full_text)
        self.assertEqual(review.current, target)
        self.assertEqual(review.current.category, "person")
        self.assertEqual(review.algorithm_signature(), signature)
        self.assertEqual(self.w.image_view.transform(), zoom)
        self.assertEqual(self.w.annotation_panel.reason_combo.currentText(), "尚未确认的原因")
        self.assertEqual(review.reason_status, "false_positive")
        self.assertEqual(review.store.get(target.id)["status"], "false_positive")
        self.w.source_panels["algorithm"].type_mapping_clear.click()
        self.assertEqual(self.w.target_table.item(0, 1).text(), "person")
        self.assertFalse(self.w.source_panels["algorithm"].type_mapping_clear.isEnabled())

    def test_source_type_mappings_are_independent_and_survive_missing_file(self):
        self.import_source(self.data, "ground_truth")
        path = self.folder / "types.txt"
        path.write_text("person:行人", encoding="utf-8")
        self.w.set_type_mapping("algorithm", TypeMapping.load(path))
        self.w.set_type_mapping("ground_truth", TypeMapping({"person": "参考行人"}))
        self.assertEqual(self.w.image_view.label_items["ground_truth"][0].full_text, "参考行人")
        self.assertEqual(self.w.target_tables["ground_truth"].item(0, 1).text(), "参考行人")
        self.w.review.all_positive()
        self.w.review.store.set_image_review("01.jpg", self.w.review.algorithm_signature(), True)
        self.w.set_type_mapping("algorithm", TypeMapping({"person": "新的显示名称"}, path))
        self.assertTrue(self.w.review.review_completed())
        self.w.close()
        path.unlink()
        self.w = MainWindow(self.settings)
        self.w.show()
        self.w.scan_folder(self.folder)
        self.wait()
        self.import_source(self.data)
        self.assertEqual(self.w.type_mappings["algorithm"].display("person"), "新的显示名称")
        self.assertEqual(self.w.type_mappings["ground_truth"].display("person"), "参考行人")
        self.assertTrue(self.w.review.review_completed())

    def test_invalid_mapping_and_cancel_keep_previous_values(self):
        self.w.set_type_mapping("algorithm", TypeMapping({"person": "行人"}))
        bad = self.folder / "bad.txt"
        bad.write_text("missing separator", encoding="utf-8")
        with patch("detretlabel.app.QFileDialog.getOpenFileName", return_value=(str(bad), "")), \
             patch("detretlabel.app.QMessageBox.warning") as warning:
            self.w.source_panels["algorithm"].type_mapping_button.click()
        warning.assert_called_once()
        with patch("detretlabel.app.QFileDialog.getOpenFileName", return_value=("", "")):
            self.w.source_panels["algorithm"].type_mapping_button.click()
        self.assertEqual(self.w.type_mappings["algorithm"].display("person"), "行人")

    def test_export_during_review_includes_hidden_images_notes_and_cause_draft(self):
        from openpyxl import load_workbook
        panel, review = self.w.annotation_panel, self.w.review
        review.choose_status("false_positive")
        target_id = review.current.id
        panel.reason_combo.setCurrentText("未确认原因草稿")
        panel.note.setText("最新备注")
        panel.note_edited.emit("最新备注")
        self.w.set_type_mapping("algorithm", TypeMapping({"person": "行人"}))
        self.w.search.setText("01")
        path = self.folder / "过程报告"
        self.assertTrue(panel.export_button.isEnabled())
        with patch("detretlabel.review_controller.QFileDialog.getSaveFileName", return_value=(str(path), "")):
            panel.export_button.click()
        self.wait()
        book = load_workbook(path.with_suffix(".xlsx"))
        try:
            self.assertEqual(book["图片进度"].max_row, 7)  # All 3 loaded images, including 2 filtered out.
            detail = book["目标明细"]
            self.assertEqual(detail["E5"].value, "行人")
            self.assertEqual(detail["K5"].value, "误检（原因待确认）")
            self.assertIsNone(detail["L5"].value)
            self.assertEqual(detail["M5"].value, "未确认原因草稿")
            self.assertEqual(detail["O5"].value, "最新备注")
        finally:
            book.close()
        self.assertEqual(review.current.id, target_id)
        self.assertEqual(review.reason_status, "false_positive")
        self.assertEqual(review.phase, "algorithm")
        self.assertEqual(panel.reason_combo.currentText(), "未确认原因草稿")
        self.assertFalse(review.review_completed())

    def test_export_cancellation_failure_and_completed_report(self):
        from openpyxl import load_workbook
        panel, review = self.w.annotation_panel, self.w.review
        path = self.folder / "report.xlsx"
        with patch("detretlabel.review_controller.QFileDialog.getSaveFileName", return_value=("", "")):
            panel.export_button.click()
        self.assertFalse(path.exists())
        path.write_bytes(b"previous report")
        with patch("detretlabel.review_controller.QFileDialog.getSaveFileName", return_value=(str(path), "")), \
             patch("detretlabel.review_controller.write_excel", side_effect=PermissionError("file in use")), \
             patch("detretlabel.app.QMessageBox.warning") as warning:
            panel.export_button.click()
            self.wait()
        warning.assert_called_once()
        self.assertEqual(path.read_bytes(), b"previous report")
        self.assertEqual(review.current.category, "person")
        self.assertTrue(panel.export_button.isEnabled())
        for _ in range(3):
            if review.phase == "algorithm":
                review.all_positive()
            review.finish_review()
        with patch("detretlabel.review_controller.QFileDialog.getSaveFileName", return_value=(str(path), "")):
            panel.export_button.click()
        self.wait()
        book = load_workbook(path)
        try:
            self.assertEqual(book["复核概览"]["B4"].value, "全部复核完成")
            self.assertEqual([book["图片进度"].cell(row, 8).value for row in range(5, 8)], ["复核完成"] * 3)
        finally:
            book.close()

    def test_clear_image_requires_confirmation_and_can_undo_complete_review(self):
        review, panel = self.w.review, self.w.annotation_panel
        review.all_positive()
        with patch("detretlabel.review_controller.QInputDialog.getText", return_value=("人工行人", True)):
            review.add_manual((1, 2, 30, 40))
        review.choose_reason("完全未检出")
        manual_id = review.current.id
        review.store.set_image_review("01.jpg", review.algorithm_signature(), True)
        other = review.role_targets("algorithm", self.folder / "02.png")[0]
        review.store.set_annotation(other, "false_positive", "其他图原因", "其他图备注")
        review.changed()
        review.save()
        initial = copy.deepcopy(review.store.records)
        history = len(review.store.history)
        with patch("detretlabel.review_controller.QMessageBox.question", return_value=QMessageBox.StandardButton.No) as confirm:
            panel.clear_image_button.click()
        self.assertEqual(confirm.call_args.args[-1], QMessageBox.StandardButton.No)
        self.assertIn("01.jpg", confirm.call_args.args[2])
        self.assertEqual(review.store.records, initial)
        self.assertEqual(review.phase, "manual")
        self.assertEqual(len(review.store.history), history)
        self.w.image_view.scale(1.5, 1.5)
        zoom = self.w.image_view.transform()
        with patch("detretlabel.review_controller.QMessageBox.question", return_value=QMessageBox.StandardButton.Yes):
            panel.clear_image_button.click()
        self.assertEqual(review.phase, "algorithm")
        self.assertEqual(self.w.current_image().name, "01.jpg")
        self.assertEqual(self.w.image_view.transform(), zoom)
        self.assertFalse(review.store.has_annotations("01.jpg"))
        self.assertEqual(review.store.get(other.id)["note"], "其他图备注")
        self.assertEqual(self.w.target_tables["manual"].rowCount(), 0)
        self.assertFalse(review.review_completed())
        self.assertFalse(panel.clear_image_button.isEnabled())
        self.assertNotIn("01.jpg", AnnotationStore.load(self.folder).image_reviews)
        self.assertEqual(len(review.store.history), history + 1)
        review.undo()
        self.assertEqual(review.store.records, initial)
        self.assertIn(manual_id, review.store.manual_targets)
        self.assertEqual(self.w.target_tables["manual"].rowCount(), 1)
        self.assertTrue(review.review_completed())

    def test_clear_all_covers_hidden_images_and_unloaded_history_with_one_undo(self):
        review, panel = self.w.review, self.w.annotation_panel
        first = review.current
        review.store.set_annotation(first, "true_positive", note="本图备注")
        hidden = review.role_targets("algorithm", self.folder / "02.png")[0]
        review.store.set_annotation(hidden, "false_positive", "隐藏图原因")
        extra = ReviewTarget.manual("not-loaded/removed.jpg", (1, 2, 30, 40))
        review.store.set_annotation(extra, "false_negative")
        review.store.set_image_review("03.jpg", review.algorithm_signature(), True)
        review.changed()
        self.w.search.setText("01")
        initial = copy.deepcopy(review.store.records)
        self.assertTrue(panel.clear_all_button.isEnabled())
        with patch("detretlabel.review_controller.QMessageBox.question", return_value=QMessageBox.StandardButton.No):
            panel.clear_all_button.click()
        self.assertEqual(review.store.records, initial)
        with patch("detretlabel.review_controller.QMessageBox.question", return_value=QMessageBox.StandardButton.Yes):
            panel.clear_all_button.click()
        self.assertFalse(review.store.has_annotations())
        self.assertFalse(AnnotationStore.load(self.folder).has_annotations())
        self.assertFalse(panel.clear_all_button.isEnabled())
        self.assertTrue(self.w.image_table.isRowHidden(1))
        review.undo()
        self.assertEqual(review.store.records, initial)
        self.assertIn(extra.id, review.store.manual_targets)
        self.assertTrue(review.store.review_completed("03.jpg", review.algorithm_signature()))

    def test_reset_flushes_pending_note_for_undo_and_preserves_disk_on_save_failure(self):
        review, panel = self.w.review, self.w.annotation_panel
        review.choose_status("false_positive")
        first = review.current.id
        review.save()
        before = review.store.path.read_bytes()
        panel.note.setText("清空前尚未自动保存的备注")
        panel.note_edited.emit(panel.note.text())
        with patch("detretlabel.review_controller.QMessageBox.question", return_value=QMessageBox.StandardButton.Yes), \
             patch("detretlabel.annotation.os.replace", side_effect=PermissionError("test")), \
             patch("detretlabel.review_controller.QMessageBox.warning") as warning:
            panel.clear_image_button.click()
        warning.assert_called_once()
        self.assertEqual(review.store.path.read_bytes(), before)
        self.assertTrue(review.store.dirty)
        self.assertFalse(review.store.records)
        review.undo()
        self.assertEqual(review.store.get(first)["note"], "清空前尚未自动保存的备注")
        self.assertEqual(review.store.get(first)["status"], "false_positive")
        self.assertTrue(review.save())

    def test_clear_can_remove_a_note_before_first_autosave(self):
        review, panel = self.w.review, self.w.annotation_panel
        target_id = review.current.id
        self.assertFalse(panel.clear_image_button.isEnabled())
        self.assertFalse(panel.clear_all_button.isEnabled())
        panel.note.setText("只有尚未保存的备注")
        panel.note_edited.emit(panel.note.text())
        self.assertTrue(panel.clear_image_button.isEnabled())
        self.assertTrue(panel.clear_all_button.isEnabled())
        with patch("detretlabel.review_controller.QMessageBox.question", return_value=QMessageBox.StandardButton.Yes):
            panel.clear_all_button.click()
        self.assertFalse(review.store.has_annotations())
        review.undo()
        self.assertEqual(review.store.get(target_id)["note"], "只有尚未保存的备注")

    def test_progress_can_move_to_another_folder_with_fresh_settings(self):
        review = self.w.review
        review.all_positive()
        with patch("detretlabel.review_controller.QInputDialog.getText", return_value=("人工目标", True)):
            review.add_manual((10, 20, 50, 60))
        manual_id = review.current.id
        review.choose_reason("人工确认漏检")
        review.finish_review()
        second_id = review.current.id
        review.store.set_annotation(review.current, None, note="交接备注")
        self.assertTrue(review.save(explicit=True))
        signature = review.algorithm_signature()
        old_source_path = self.w.sources["algorithm"].path
        self.w.close()
        transfer = tempfile.TemporaryDirectory()
        self.addCleanup(transfer.cleanup)
        destination = Path(transfer.name).resolve() / "copied-images"
        shutil.copytree(self.folder, destination)
        self.folder = destination
        fresh = QSettings(str(Path(transfer.name) / "new-machine.ini"), QSettings.Format.IniFormat)
        self.w = MainWindow(fresh)
        self.w.show()
        self.w.scan_folder(destination)
        self.wait()
        self.import_source(self.data)
        resumed = self.w.review
        self.assertNotEqual(self.w.sources["algorithm"].path, old_source_path)
        self.assertEqual(resumed.algorithm_signature(), signature)
        self.assertTrue(resumed.review_completed(destination / "01.jpg"))
        self.assertEqual(resumed.store.get(manual_id)["reason"], "人工确认漏检")
        self.assertIn(manual_id, resumed.store.manual_targets)
        self.assertEqual(resumed.store.get(second_id)["note"], "交接备注")
        self.assertIsNone(resumed.store.get(second_id)["status"])
        resumed.next_review_image()
        self.assertEqual(self.w.current_image().name, "02.png")
        self.assertEqual(resumed.current.id, second_id)

    def test_finishing_does_not_advance_if_save_fails(self):
        review = self.w.review
        review.all_positive()
        with patch("detretlabel.annotation.os.replace", side_effect=PermissionError("test")), \
             patch("detretlabel.review_controller.QMessageBox.warning"):
            review.finish_review()
        self.assertEqual(self.w.current_image().name, "01.jpg")
        self.assertEqual(review.phase, "manual")
        self.assertTrue(review.store.dirty)
        review.finish_review()
        self.assertEqual(self.w.current_image().name, "02.png")


if __name__ == "__main__":
    unittest.main()
