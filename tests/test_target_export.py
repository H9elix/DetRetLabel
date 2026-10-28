import json
import re
import tempfile
import unittest
from dataclasses import replace
from decimal import Decimal
from pathlib import Path
from unittest.mock import Mock, patch
from openpyxl import load_workbook

from detretlabel.annotation import AnnotationStore, ReviewTarget
from detretlabel.core import Collection, guess_mapping, import_results
from detretlabel.review_report import build_report
from detretlabel.target_export import (
    OUTPUT_FIELDS, ExportIssue, build_target_export, check_report_path, write_export_files, write_targets_json,
)
from detretlabel.type_mapping import TypeMapping


class TargetExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.images = [self.folder / "sub/one.jpg", self.folder / "two.png", self.folder / "empty.jpg"]
        self.a = ReviewTarget("a", "sub/one.jpg", "algorithm", (20, 10, 100, 60), 7, .95)
        self.b = replace(self.a, id="b", category="误检")
        self.c = replace(self.a, id="c", image="two.png", category="未审核", score=None)
        self.manual = ReviewTarget.manual("sub/one.jpg", (40, 20, 80, 50), "人工行人")
        self.manual_two = ReviewTarget.manual("two.png", (10, 5, 30, 40), "car")
        self.pending = ReviewTarget.manual("two.png", (1, 2, 3, 4))
        self.by_image = {self.images[0]: [self.a, self.b, self.manual, replace(self.a, id="gt", role="ground_truth")],
                         self.images[1]: [self.c, self.manual_two, self.pending]}
        self.by_image[self.images[0]].extend([
            replace(self.a, id="known-manual-person", category="人工行人"),
            replace(self.a, id="known-manual-car", category="car")])
        self.store = AnnotationStore(self.folder)
        self.store.set_annotation(self.a, "true_positive", note="=中文正检备注")
        self.store.set_annotation(self.b, "false_positive", "种类检错")
        self.store.set_annotation(self.manual, "false_negative", "未检出", "补漏备注")
        self.store.set_annotation(self.manual_two, "false_negative")
        self.store.set_annotation(self.pending, None)
        self.sizes = {"sub/one.jpg": (200, 100), "two.png": (200, 100)}
        self.mapping = TypeMapping({"7": "1", "人工行人": "行人", "car": "汽车"})

    def tearDown(self):
        self.temp.cleanup()

    def snapshot(self):
        return build_report(self.folder, self.images, self.by_image, self.store, "sig", {}, {}, {})

    def export(self, **kwargs):
        return build_target_export(self.snapshot(), image_sizes=self.sizes, type_mapping=self.mapping, **kwargs)

    def test_fixed_schema_scope_normalization_and_numeric_string_names(self):
        export = self.export()
        self.assertFalse(export.issues)
        self.assertEqual(export.candidate_count, 3)
        self.assertEqual([target["name"] for target in export.targets], ["sub/one.jpg", "sub/one.jpg", "two.png"])
        target = export.targets[0]
        self.assertEqual(tuple(target), OUTPUT_FIELDS)
        self.assertEqual(target, dict(name="sub/one.jpg", width=200, height=100, type=7, type_name="1",
            conf=Decimal(".95000000"), x=Decimal(".10000000"), y=Decimal(".10000000"),
            w=Decimal(".40000000"), h=Decimal(".50000000")))
        self.assertEqual(export.targets[1]["conf"], 1)
        self.assertEqual(export.targets[1]["type_name"], "行人")
        self.assertTrue(all(tuple(row) == OUTPUT_FIELDS and not any(isinstance(v, (dict, list)) for v in row.values()) for row in export.targets))

    def test_decimal_writer_has_eight_digits_no_scientific_notation_and_preserves_type(self):
        report = self.snapshot()
        first = report.targets[0]
        extra = dict(first, category="7", target_id="string", xmin=.000002)
        report.targets.append(extra)
        report.targets.append(dict(first, category=1e-12, target_id="tiny-type"))
        export = build_target_export(report, image_sizes=self.sizes, type_mapping=self.mapping)
        path = write_targets_json(export, self.folder / "结果.json")
        text = path.read_text(encoding="utf-8")
        decoded = json.loads(text)
        self.assertIsInstance(decoded[0]["type"], int)
        self.assertEqual(decoded[3]["type"], "7")
        self.assertIsInstance(decoded[3]["type_name"], str)
        self.assertEqual(decoded[4]["type"], 1e-12)  # Category identifiers must never be rounded.
        self.assertIn('"type": 0.000000000001', text)
        self.assertIn('"x": 0.00000001', text)
        self.assertIn('"conf": 1.00000000', text)
        self.assertFalse(re.search(r':\s*-?\d+(?:\.\d+)?[eE][+-]?\d+', text))
        values = re.findall(r'"(?:conf|x|y|w|h)":\s*(-?\d+\.\d+)', text)
        self.assertTrue(all(len(value.split('.')[1]) == 8 for value in values))

    def test_missing_conf_and_mapping_use_null_and_string_fallback_with_warnings(self):
        self.store.set_annotation(self.c, "true_positive")
        export = self.export()
        target = next(target for target in export.targets if target["type"] == "未审核")
        self.assertIsNone(target["conf"])
        self.assertEqual(target["type_name"], "未审核")
        self.assertEqual({issue.code for issue in export.warnings}, {"missing_conf", "type_name"})
        files = write_export_files(export, self.folder / "null.json")
        self.assertIn('"conf": null', files.json.read_text(encoding="utf-8"))
        book = load_workbook(files.check_report)
        try:
            self.assertEqual(book.sheetnames, ["检查明细", "导出概览"])
            sheet = book["检查明细"]
            self.assertEqual(sheet.tables["ExportChecks"].autoFilter.ref, "A1:J3")
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertIsNone(book["导出概览"].freeze_panes)
            self.assertTrue(any("JSON null" in str(row[9].value) for row in list(sheet.rows)[1:]))
        finally:
            book.close()

    def test_all_input_shapes_and_coordinate_formats_export_the_same_schema(self):
        target = {"bbox": [20, 10, 100, 60], "type": "car", "conf": .8}
        cases = [([{**target, "name": "one.jpg"}], Collection("")),
                 ({"results": [{"path": "one.jpg", "detections": [target]}]}, Collection("/results")),
                 ({"one.jpg": [target]}, Collection("", "keyed")),
                 ([{"name": "one.jpg", "geometry": {"bbox": target["bbox"]}, "semantic": {"type": "car"}, "conf": .8}], Collection("")),
                 ([{"name": "one.jpg", "x": .1, "y": .1, "w": .4, "h": .5, "type": "car", "conf": .8}], Collection(""))]
        outputs = []
        for index, (document, collection) in enumerate(cases):
            mapping = guess_mapping(document, collection)
            if index == 4:
                mapping = replace(mapping, normalized=True)
            result = import_results(document, mapping, self.images, image_size_loader=lambda image: (200, 100))
            self.assertFalse(result.issues)
            detection = result.detections[self.images[0]][0]
            target = ReviewTarget.from_detection("sub/one.jpg", "algorithm", None, "sig", detection)
            store = AnnotationStore()
            store.set_annotation(target, "true_positive")
            report = build_report(self.folder, self.images, {self.images[0]: [target]}, store, "sig", {}, {}, {})
            export = build_target_export(report, image_sizes=self.sizes, type_mapping=self.mapping)
            path = write_targets_json(export, self.folder / f"input-{index}.json")
            outputs.append(json.loads(path.read_text(encoding="utf-8")))
        self.assertTrue(all(output == outputs[0] for output in outputs))

    def test_actual_dimensions_win_and_each_image_is_read_only_once(self):
        loader = Mock(return_value=(400, 200))
        export = self.export(image_size_loader=loader)
        self.assertEqual(loader.call_count, 2)
        self.assertEqual(export.targets[0]["width"], 400)
        self.assertEqual(export.targets[0]["x"], Decimal(".05000000"))
        self.assertEqual({issue.code for issue in export.warnings}, {"size_conflict"})

    def test_saved_dimensions_support_export_without_images_and_reject_conflicting_metadata(self):
        export = self.export(image_size_loader=Mock(side_effect=OSError("offline")))
        self.assertFalse(export.errors)
        conflicting = self.export(size_issues={"sub/one.jpg": "02 中同图宽高不一致"})
        self.assertTrue(conflicting.errors)
        self.assertEqual({issue.code for issue in conflicting.errors}, {"dimensions"})
        confirmed = self.export(size_issues={"sub/one.jpg": "02 中同图宽高不一致"}, trusted_sizes={"sub/one.jpg"})
        self.assertFalse(confirmed.errors)
        self.assertIn("size_mapping", {issue.code for issue in confirmed.warnings})

    def test_missing_dimensions_blocks_only_json_and_writes_report_without_overwriting(self):
        export = build_target_export(self.snapshot(), type_mapping=self.mapping)
        path = self.folder / "existing.json"
        path.write_bytes(b"old result")
        files = write_export_files(export, path)
        self.assertIsNone(files.json)
        self.assertEqual(path.read_bytes(), b"old result")
        book = load_workbook(files.check_report)
        try:
            self.assertIn("JSON 未生成", book["导出概览"]["B2"].value)
            self.assertTrue(all(row[8].value == "未生成 JSON" for row in list(book["检查明细"].rows)[1:]))
        finally:
            book.close()
        self.assertEqual({issue.image for issue in export.errors}, {"sub/one.jpg", "two.png"})

    def test_partial_output_requires_explicit_choice_and_reports_skipped_targets(self):
        export = build_target_export(self.snapshot(), image_sizes={"sub/one.jpg": (200, 100)}, type_mapping=self.mapping)
        with self.assertRaises(ValueError):
            write_targets_json(export, self.folder / "not-allowed.json")
        files = write_export_files(export, self.folder / "partial.json", allow_partial=True)
        self.assertEqual((files.target_count, files.skipped_count), (2, 1))
        self.assertEqual(len(json.loads(files.json.read_text(encoding="utf-8"))), 2)
        book = load_workbook(files.check_report)
        try:
            self.assertEqual(book["导出概览"]["B2"].value, "仅导出有效部分")
            self.assertTrue(any(row[2].value == "two.png" and row[8].value == "跳过该目标" for row in list(book["检查明细"].rows)[1:]))
        finally:
            book.close()

    def test_missing_type_invalid_numbers_and_zero_after_rounding_are_blocking(self):
        for change, code in (({"category": None}, "type"), ({"score": float("nan")}, "conf"),
                             ({"xmax": 20 + 1e-10}, "precision"), ({"ymax": float("inf")}, "box")):
            with self.subTest(change=change):
                report = self.snapshot()
                report.targets[0].update(change)
                export = build_target_export(report, image_sizes=self.sizes, type_mapping=self.mapping)
                self.assertIn(code, {issue.code for issue in export.errors})

    def test_outside_boxes_are_kept_and_reported_without_clipping(self):
        report = self.snapshot()
        report.targets[0]["xmin"] = -20
        report.targets[0]["xmax"] = 220
        export = build_target_export(report, image_sizes=self.sizes, type_mapping=self.mapping)
        self.assertFalse(export.errors)
        self.assertEqual((export.targets[0]["x"], export.targets[0]["w"]), (Decimal("-.10000000"), Decimal("1.20000000")))
        self.assertIn("outside_image", {issue.code for issue in export.warnings})

    def test_empty_output_is_valid_and_previous_diagnostic_is_refreshed(self):
        report = replace(self.snapshot(), targets=[])
        export = build_target_export(report)
        path = self.folder / "empty.json"
        check_report_path(path).write_text("old errors", encoding="utf-8")
        files = write_export_files(export, path)
        self.assertEqual(json.loads(path.read_text(encoding="utf-8")), [])
        book = load_workbook(files.check_report)
        try:
            self.assertEqual(book["检查明细"].max_row, 1)
            self.assertEqual(book["检查明细"].auto_filter.ref, "A1:J1")
            self.assertEqual(book["导出概览"]["B2"].value, "导出完成")
        finally:
            book.close()

    def test_legacy_manual_types_are_preserved_and_filterable_in_excel(self):
        report = self.snapshot()
        manual = next(row for row in report.targets if row["target_id"] == self.manual.id)
        manual["category"] = "7"
        export = build_target_export(report, image_sizes=self.sizes, type_mapping=self.mapping)
        self.assertFalse(export.errors)
        target = next(target for target in export.targets if target["type"] == "7")
        self.assertIsInstance(target["type"], str)
        self.assertEqual(target["type_name"], "1")
        files = write_export_files(export, self.folder / "legacy.json")
        book = load_workbook(files.check_report)
        try:
            rows = list(book["检查明细"].values)
            row = next(row for row in rows[1:] if row[1] == "人工种类与算法类别不一致")
            self.assertEqual((row[5], row[6], row[7], row[8]), ('"7"', "字符串", "1", "保留并导出"))
            self.assertIn("不会根据显示名称反推", row[9])
        finally:
            book.close()

    def test_second_replace_failure_restores_existing_json_and_report(self):
        from detretlabel import target_export
        real_replace = target_export.os.replace
        export = build_target_export(self.snapshot(), image_sizes=self.sizes)
        path = self.folder / "atomic.json"
        report_path = check_report_path(path)
        path.write_bytes(b"old json")
        report_path.write_bytes(b"old check")
        def fail_report(source, destination):
            if Path(destination) == report_path:
                raise PermissionError("report locked")
            return real_replace(source, destination)
        with patch("detretlabel.target_export.os.replace", side_effect=fail_report), self.assertRaises(PermissionError):
            write_export_files(export, path)
        self.assertEqual(path.read_bytes(), b"old json")
        self.assertEqual(report_path.read_bytes(), b"old check")
        self.assertFalse(list(self.folder.glob(".detretlabel-export-*")))

    def test_nonfinite_mutation_never_overwrites_a_previous_json(self):
        export = self.export()
        export.targets[0]["conf"] = Decimal("NaN")
        path = self.folder / "existing.json"
        path.write_bytes(b"old json")
        with self.assertRaises(ValueError):
            write_targets_json(export, path)
        self.assertEqual(path.read_bytes(), b"old json")
        self.assertFalse(list(self.folder.glob("review-targets-*.json")))

    def test_excel_check_values_are_literal_text_not_formulas(self):
        export = self.export()
        export.issues.append(ExportIssue("warning", "=image.jpg", "=id", "custom", "=1+1", "=label", "=type", "=name"))
        files = write_export_files(export, self.folder / "literal.json")
        book = load_workbook(files.check_report)
        try:
            self.assertTrue(all(cell.data_type != "f" for sheet in book for row in sheet for cell in row))
            self.assertEqual(book["检查明细"]["J2"].value, "=1+1")
        finally:
            book.close()


if __name__ == "__main__":
    unittest.main()
