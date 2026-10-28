import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from detretlabel.annotation import AnnotationStore, ReviewTarget
from detretlabel.review_report import build_report, write_excel
from detretlabel.type_mapping import TypeMapping


class ReviewReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.images = [self.folder / "one.jpg", self.folder / "two.png", self.folder / "empty.jpg"]
        self.a = ReviewTarget("a", "one.jpg", "algorithm", (1.25, 2, 31, 42), "001", .912345, "alg.json")
        self.b = replace(self.a, id="b", category="person")
        self.c = replace(self.a, id="c", image="two.png", category=0, score=90)
        self.gt = replace(self.a, id="gt", role="ground_truth")
        self.manual = ReviewTarget.manual("one.jpg", (10, 20, 30, 40), "汽车")
        self.pending = ReviewTarget.manual("two.png", (11, 22, 33, 44))
        self.by_image = {self.images[0]: [self.a, self.b, self.manual, self.gt], self.images[1]: [self.c, self.pending]}
        self.store = AnnotationStore(self.folder)
        self.store.set_annotation(self.a, "true_positive", note="=中文备注")
        self.store.set_annotation(self.b, "false_positive", "种类检错", "保留备注")
        self.store.set_annotation(self.manual, "false_negative", "完全未检出")
        self.store.set_annotation(self.pending, None)
        self.store.set_image_review("one.jpg", "sig", True)
        self.store.set_image_review("empty.jpg", "sig", True)

    def tearDown(self):
        self.temp.cleanup()

    def snapshot(self, **kwargs):
        return build_report(self.folder, self.images, self.by_image, self.store, "sig",
            {"algorithm": TypeMapping({"001": "行人", "0": "汽车"}), "ground_truth": TypeMapping()},
            {"algorithm": "alg.json", "ground_truth": "gt.json"},
            {self.images[2]: "空目标列表"}, **kwargs)

    def table_rows(self, sheet, name):
        left, top, right, bottom = range_boundaries(sheet.tables[name].ref)
        headers = [sheet.cell(top, column).value for column in range(left, right + 1)]
        return [dict(zip(headers, [sheet.cell(row, column).value for column in range(left, right + 1)]))
                for row in range(top + 1, bottom + 1)]

    def test_separate_result_sheets_have_matching_totals_categories_and_full_target_lists(self):
        report = self.snapshot()
        book = load_workbook(write_excel(report, self.folder / "classified.xlsx"))
        try:
            for name, key, prefix, expected_id in (
                    ("正检", "true_positive", "TruePositive", self.a.id),
                    ("误检", "false_positive", "FalsePositive", self.b.id),
                    ("漏检", "false_negative", "FalseNegative", self.manual.id)):
                sheet = book[name]
                self.assertEqual(sheet["B4"].value, report.stats[key])
                self.assertEqual(sheet["E4"].value, 1)
                self.assertEqual(sheet["H4"].value, 1)
                categories = self.table_rows(sheet, prefix + "Categories")
                targets = self.table_rows(sheet, prefix + "Targets")
                self.assertEqual(sum(row[f"{name}数量"] for row in categories), sheet["B4"].value)
                self.assertEqual(len(targets), sheet["B4"].value)
                self.assertEqual(targets[0]["目标唯一标识"], expected_id)
                self.assertEqual(targets[0]["图片相对路径"], "one.jpg")
                self.assertIn("xmin", targets[0])
                self.assertIn("备注", targets[0])
                self.assertIn("原因（已保存）", targets[0])
                self.assertIsNone(sheet.freeze_panes)
                self.assertEqual(sheet.tables[prefix + "Targets"].autoFilter.ref, sheet.tables[prefix + "Targets"].ref)
                _, header, _, _ = range_boundaries(sheet.tables[prefix + "Targets"].ref)
                self.assertEqual(sheet["B5"].hyperlink.location, f"'{name}'!A{header}")
            positive = self.table_rows(book["正检"], "TruePositiveTargets")[0]
            self.assertEqual(positive["备注"], "=中文备注")
            self.assertEqual(positive["置信度（原始尺度）"], .912345)
            self.assertEqual(positive["显示种类"], "行人")
        finally:
            book.close()

    def test_per_class_counts_distinguish_numeric_and_string_types_with_the_same_display_name(self):
        self.store.set_annotation(self.c, "true_positive")
        for target in (replace(self.a, id="numeric-zero", category=0), replace(self.a, id="string-zero", category="0")):
            self.by_image[self.images[0]].append(target)
            self.store.set_annotation(target, "true_positive")
        book = load_workbook(write_excel(self.snapshot(), self.folder / "types.xlsx"))
        try:
            sheet = book["正检"]
            self.assertEqual(sheet["B4"].value, 4)
            self.assertEqual(sheet["H4"].value, 3)
            rows = self.table_rows(sheet, "TruePositiveCategories")
            counts = {(row["原始种类"], row["type 值类型"]): row["正检数量"] for row in rows}
            self.assertEqual(counts, {("0", "数值"): 2, ("0", "字符串"): 1, ("001", "字符串"): 1})
            self.assertEqual({row["显示种类"] for row in rows if row["原始种类"] == "0"}, {"汽车"})
        finally:
            book.close()

    def test_pending_reasons_and_unknown_manual_types_are_visible_without_including_old_algorithm_misses(self):
        self.store.set_annotation(self.a, "false_negative", "旧算法漏检")
        self.store.set_annotation(self.b, "false_positive", reason_pending=True, reason_draft="待确认原因")
        self.store.set_annotation(self.pending, "false_negative")
        book = load_workbook(write_excel(self.snapshot(), self.folder / "pending.xlsx"))
        try:
            self.assertEqual(book["正检"]["B4"].value, 0)
            sheet = book["误检"]
            self.assertEqual(sheet["B4"].value, 1)
            self.assertEqual(sheet["E5"].value, 1)
            self.assertEqual(self.table_rows(sheet, "FalsePositiveCategories")[0]["其中原因待确认"], 1)
            row = self.table_rows(sheet, "FalsePositiveTargets")[0]
            self.assertEqual(row["人工结论"], "误检（原因待确认）")
            self.assertEqual(row["原因草稿（未确认）"], "待确认原因")
            sheet = book["漏检"]
            self.assertEqual(sheet["B4"].value, 2)
            categories = self.table_rows(sheet, "FalseNegativeCategories")
            self.assertIn("未指定种类", [row["原始种类"] for row in categories])
            ids = {row["目标唯一标识"] for row in self.table_rows(sheet, "FalseNegativeTargets")}
            self.assertEqual(ids, {self.manual.id, self.pending.id})
            self.assertNotIn(self.a.id, ids)
        finally:
            book.close()

    def test_result_sheets_include_decisions_on_unfinished_images(self):
        self.store.set_annotation(self.c, "true_positive")
        book = load_workbook(write_excel(self.snapshot(), self.folder / "in-progress.xlsx"))
        try:
            sheet = book["正检"]
            self.assertEqual(sheet["B4"].value, 2)
            self.assertEqual(sheet["E4"].value, 2)
            self.assertEqual(sheet["H5"].value, 1)
            self.assertEqual({row["图片相对路径"] for row in self.table_rows(sheet, "TruePositiveTargets")}, {"one.jpg", "two.png"})
        finally:
            book.close()

    def test_scope_statistics_completion_and_draft_are_independent_of_gt(self):
        old = replace(self.a, id="old-version", signature="old")
        self.store.set_annotation(old, "false_positive", "旧记录")
        self.store.set_image_review("one.jpg", "sig", True)
        report = self.snapshot(pending_reason=("b", "新原因草稿"))
        self.assertEqual(report.stats, dict(images=3, completed_images=2, pending_images=1,
            algorithm=3, true_positive=1, false_positive=1, unreviewed=1, false_negative=1, pending_manual=1,
            pending_reasons=0, algorithm_completed_images=2, miss_completed_images=2))
        self.assertEqual(len(report.targets), 5)
        self.assertEqual(report.images[0]["gt_count"], 1)
        self.assertEqual(report.images[0]["review_state"], "复核完成")
        self.assertEqual(report.images[2]["input_state"], "空目标列表")
        row = next(row for row in report.targets if row["target_id"] == "b")
        self.assertEqual((row["reason"], row["draft"], row["editing"]), ("种类检错", "新原因草稿", "正在选择原因"))
        self.assertEqual(report.targets[0]["raw_type"], "001")
        self.assertEqual(report.targets[0]["display_type"], "行人")
        self.assertEqual(sum(r["false_negative"] for r in report.categories), 1)
        self.store.set_annotation(self.b, "true_positive")
        self.assertEqual(row["status"], "误检")  # Export snapshot does not follow later UI edits.

    def test_legacy_algorithm_misses_and_mismatched_signatures_remain_pending(self):
        self.store.set_annotation(self.a, "false_negative", "旧数据")
        self.store.set_image_review("one.jpg", "sig", True)
        self.store.set_image_review("empty.jpg", "other", True)
        report = self.snapshot()
        self.assertEqual(report.stats["completed_images"], 0)
        self.assertEqual(report.stats["unreviewed"], 2)
        self.assertEqual(report.targets[0]["status"], "旧漏检，待重审")
        self.assertEqual(report.stats["false_negative"], 1)

    def test_excel_contains_all_sheets_literal_text_coordinates_and_filters(self):
        path = write_excel(self.snapshot(pending_reason=("b", "临时输入")), self.folder / "报告.xlsx")
        book = load_workbook(path)
        try:
            self.assertEqual(book.sheetnames, ["复核概览", "图片进度", "目标明细", "正检", "误检", "漏检"])
            sheet = book["目标明细"]
            self.assertEqual(sheet["D5"].value, "001")
            self.assertEqual(sheet["D5"].data_type, "s")
            self.assertEqual(sheet["E5"].value, "行人")
            self.assertEqual(sheet["F5"].value, .912345)
            self.assertEqual(sheet["G5"].value, 1.25)
            self.assertEqual(sheet["O5"].value, "=中文备注")
            self.assertEqual(sheet["O5"].data_type, "s")
            self.assertEqual(sheet["M6"].value, "临时输入")
            self.assertEqual(sheet.freeze_panes, "D5")
            self.assertEqual(sheet.tables["TargetReviewDetails"].ref, "A4:U9")
            self.assertEqual(book["图片进度"].max_row, 7)
            self.assertEqual(book["复核概览"]["B4"].value, "复核进行中")
            self.assertIsNone(book["复核概览"].freeze_panes)
            self.assertEqual(book["图片进度"].freeze_panes, "D5")
            self.assertFalse(any(c.data_type == "f" for s in book for row in s for c in row))
        finally:
            book.close()

    def test_empty_results_still_export_headers_and_image_progress(self):
        report = build_report(self.folder, self.images, {}, AnnotationStore(), "sig", {}, {}, {})
        book = load_workbook(write_excel(report, self.folder / "empty.xlsx"))
        try:
            self.assertEqual(book["图片进度"].max_row, 7)
            self.assertEqual(book["目标明细"].max_row, 4)
            self.assertEqual(report.stats["pending_images"], 3)
            for name in ("正检", "误检", "漏检"):
                self.assertEqual(book[name]["B4"].value, 0)
                self.assertEqual(book[name]["H4"].value, 0)
                self.assertEqual(book[name]["A9"].value, f"暂无{name}目标")
                self.assertEqual(book[name].max_row, 12)
                self.assertIsNone(book[name].freeze_panes)
        finally:
            book.close()

    def test_partial_algorithm_and_saved_reason_draft_are_distinct_from_completed_miss_review(self):
        self.store.set_annotation(self.b, "false_positive", note="仍需确认原因", reason_pending=True, reason_draft="暂定种类错误")
        self.store.set_image_review("one.jpg", "sig", True)
        report = self.snapshot()
        row = report.images[0]
        self.assertEqual(row["algorithm_state"], "未完成")
        self.assertEqual(row["miss_review_state"], "已完成")
        self.assertEqual(row["review_state"], "03未完成 / 04已完成")
        self.assertTrue(row["miss_completed_at"])
        self.assertFalse(row["completed_at"])
        self.assertEqual(report.stats["completed_images"], 1)
        self.assertEqual(report.stats["miss_completed_images"], 2)
        self.assertEqual(report.stats["pending_reasons"], 1)
        detail = next(r for r in report.targets if r["target_id"] == "b")
        self.assertEqual(detail["draft"], "暂定种类错误")
        self.assertEqual(detail["editing"], "原因待确认")
        book = load_workbook(write_excel(report, self.folder / "partial.xlsx"))
        try:
            progress = book["图片进度"]
            self.assertEqual(progress["H5"].value, "03未完成 / 04已完成")
            self.assertEqual(progress["L5"].value, "未完成")
            self.assertEqual(progress["M5"].value, 1)
            self.assertEqual(progress["N5"].value, "已完成")
            self.assertTrue(progress["O5"].value)
            self.assertEqual(book["目标明细"]["M6"].value, "暂定种类错误")
            self.assertEqual(book["目标明细"]["T6"].value, "未完成")
            self.assertEqual(book["目标明细"]["U6"].value, "已完成")
        finally:
            book.close()

    def test_atomic_failure_preserves_existing_report_and_cleans_temporary_file(self):
        path = self.folder / "existing.xlsx"
        path.write_bytes(b"old report")
        with patch("detretlabel.review_report.os.replace", side_effect=PermissionError("Excel is open")):
            with self.assertRaises(PermissionError):
                write_excel(self.snapshot(), path)
        self.assertEqual(path.read_bytes(), b"old report")
        self.assertFalse(list(self.folder.glob("review-report-*.xlsx")))

    def test_invalid_cell_data_is_reported_without_overwriting(self):
        path = self.folder / "existing.xlsx"
        path.write_bytes(b"old report")
        for invalid in ("text\x00", "长" * 32768):
            with self.subTest(invalid=invalid[:20]):
                self.store.set_annotation(self.a, "true_positive", note=invalid)
                with self.assertRaises(ValueError):
                    write_excel(self.snapshot(), path)
                self.assertEqual(path.read_bytes(), b"old report")
        report = self.snapshot()
        report.targets[0]["score"] = float("nan")
        with self.assertRaises(ValueError):
            write_excel(report, path)


if __name__ == "__main__":
    unittest.main()
